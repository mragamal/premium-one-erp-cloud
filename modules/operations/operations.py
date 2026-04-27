import csv
import io
import os
import uuid
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from urllib.parse import quote

from fastapi import APIRouter, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse

from audit import actor_name_from_request, render_audit_log_card, safe_log_action
from db import get_conn
from layout import render_page

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

router = APIRouter()


SERVICE_CATEGORIES = [
    ("field_maintenance", "Field Maintenance"),
    ("workshop_repair", "Workshop Repair"),
    ("cabinet_assembly", "Cabinet Assembly"),
]

ZONE_LEVELS = [
    ("zone_1", "Zone 1"),
    ("zone_2", "Zone 2"),
    ("zone_3", "Zone 3"),
    ("zone_4", "Zone 4"),
]

WORK_ORDER_STATUSES = [
    ("new", "New"),
    ("assigned", "Assigned"),
    ("in_progress", "In Progress"),
    ("submitted", "Submitted"),
    ("approved", "Approved"),
    ("rejected", "Rejected"),
    ("closed", "Closed"),
]

TRIP_STATUSES = [
    ("draft", "Draft"),
    ("dispatched", "Dispatched"),
    ("completed", "Completed"),
    ("approved", "Approved"),
]

DRIVER_SOURCE_OPTIONS = [
    ("company_driver", "Company Driver"),
    ("supplier_driver", "Supplier Driver"),
]


def safe(value):
    return "" if value is None else str(value).strip()


def safe_int(value, default=0):
    try:
        if value is None or str(value).strip() == "":
            return default
        return int(float(value))
    except Exception:
        return default


def to_decimal(value, default="0"):
    try:
        text = safe(value).replace(",", "")
        if text in ["", ".", "-", "-."]:
            text = default
        return Decimal(text)
    except (InvalidOperation, ValueError, TypeError):
        return Decimal(default)


def q2(value):
    return to_decimal(value).quantize(Decimal("1.00"), rounding=ROUND_HALF_UP)


def money(value):
    return f"{q2(value):,.2f}"


def ensure_column(conn, table_name, column_name, alter_sql):
    cols = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    names = [c["name"] for c in cols]
    if column_name not in names:
        conn.execute(alter_sql)


def ensure_tables():
    conn = get_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_contract_companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            name TEXT NOT NULL,
            contact_person TEXT,
            phone TEXT,
            email TEXT,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_contracts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            contract_no TEXT,
            company_id INTEGER NOT NULL,
            contract_name TEXT NOT NULL,
            pricing_method TEXT DEFAULT 'standard',
            start_date TEXT,
            end_date TEXT,
            status TEXT DEFAULT 'active',
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_rental_suppliers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            name TEXT NOT NULL,
            contact_person TEXT,
            phone TEXT,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_fault_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            name TEXT NOT NULL,
            service_category TEXT,
            default_service_price REAL DEFAULT 0,
            default_incentive REAL DEFAULT 0,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_regions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            name TEXT NOT NULL,
            zone_level TEXT,
            allowance_amount REAL DEFAULT 0,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_service_catalog (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            name TEXT NOT NULL,
            service_category TEXT,
            unit_price REAL DEFAULT 0,
            technician_incentive REAL DEFAULT 0,
            default_region_level TEXT,
            default_duration_hours REAL DEFAULT 0,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_vehicle_rates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            rental_supplier_id INTEGER,
            slab_name TEXT,
            vehicle_type TEXT NOT NULL,
            ticket_open_price REAL DEFAULT 0,
            second_slab_to_km REAL DEFAULT 300,
            km_rate_101_300 REAL DEFAULT 0,
            km_rate_over_300 REAL DEFAULT 0,
            km_rate REAL DEFAULT 0,
            waiting_hour_rate REAL DEFAULT 0,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_vehicles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            vehicle_name TEXT,
            vehicle_type TEXT NOT NULL,
            rental_supplier_id INTEGER,
            vehicle_rate_id INTEGER,
            pricing_slab_name TEXT,
            plate_no TEXT,
            driver_source TEXT DEFAULT 'company_driver',
            supplier_driver_name TEXT,
            notes TEXT,
            is_active INTEGER DEFAULT 1,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_action_price_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            contract_id INTEGER NOT NULL,
            action_id INTEGER NOT NULL,
            version_name TEXT,
            effective_from TEXT NOT NULL,
            effective_to TEXT,
            fuel_reference TEXT,
            action_price REAL DEFAULT 0,
            technician_incentive REAL DEFAULT 0,
            region_allowance REAL DEFAULT 0,
            is_active INTEGER DEFAULT 1,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_vehicle_price_versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            contract_id INTEGER NOT NULL,
            vehicle_rate_id INTEGER NOT NULL,
            version_name TEXT,
            effective_from TEXT NOT NULL,
            effective_to TEXT,
            fuel_reference TEXT,
            ticket_open_price REAL DEFAULT 0,
            second_slab_to_km REAL DEFAULT 300,
            km_rate_101_300 REAL DEFAULT 0,
            km_rate_over_300 REAL DEFAULT 0,
            waiting_hour_rate REAL DEFAULT 0,
            is_active INTEGER DEFAULT 1,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_action_materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            action_id INTEGER NOT NULL,
            line_no INTEGER DEFAULT 1,
            item_id INTEGER,
            item_code TEXT,
            item_name TEXT,
            uom TEXT,
            qty REAL DEFAULT 0,
            unit_cost REAL DEFAULT 0,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_work_orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            work_order_no TEXT,
            company_id INTEGER,
            fault_type_id INTEGER,
            service_id INTEGER,
            region_id INTEGER,
            request_date TEXT,
            site_code TEXT,
            site_name TEXT,
            complaint_details TEXT,
            required_materials TEXT,
            technician_id INTEGER,
            manager_id INTEGER,
            priority TEXT DEFAULT 'normal',
            trip_required INTEGER DEFAULT 1,
            service_price REAL DEFAULT 0,
            technician_incentive REAL DEFAULT 0,
            region_allowance REAL DEFAULT 0,
            status TEXT DEFAULT 'new',
            created_by TEXT,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT,
            closure_notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_tickets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ticket_no TEXT,
            company_id INTEGER NOT NULL,
            contract_id INTEGER,
            ticket_date TEXT,
            fault_type_id INTEGER,
            site_code TEXT,
            site_name TEXT,
            priority TEXT DEFAULT 'normal',
            request_channel TEXT,
            complaint_details TEXT,
            status TEXT DEFAULT 'open',
            created_by TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_trip_tickets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trip_no TEXT,
            trip_date TEXT,
            vehicle_id INTEGER NOT NULL,
            rental_supplier_id INTEGER,
            vehicle_type TEXT,
            vehicle_rate_id INTEGER,
            driver_source TEXT DEFAULT 'company_driver',
            driver_employee_id INTEGER,
            supplier_driver_name TEXT,
            status TEXT DEFAULT 'draft',
            start_odometer REAL DEFAULT 0,
            start_photo_path TEXT,
            end_odometer REAL DEFAULT 0,
            end_photo_path TEXT,
            waiting_hours REAL DEFAULT 0,
            total_km REAL DEFAULT 0,
            ticket_open_price REAL DEFAULT 0,
            second_slab_to_km REAL DEFAULT 300,
            km_rate_101_300 REAL DEFAULT 0,
            km_rate_over_300 REAL DEFAULT 0,
            waiting_hour_rate REAL DEFAULT 0,
            total_cost REAL DEFAULT 0,
            allocated_cost_per_work_order REAL DEFAULT 0,
            driver_commission_pct REAL DEFAULT 0,
            driver_commission_amount REAL DEFAULT 0,
            notes TEXT,
            movement_notes TEXT,
            accounting_notes TEXT,
            created_by TEXT,
            movement_closed_by TEXT,
            approved_by TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT,
            approved_at TEXT
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_trip_work_orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trip_id INTEGER NOT NULL,
            work_order_id INTEGER NOT NULL,
            line_no INTEGER DEFAULT 1,
            allocated_cost REAL DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ops_technician_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            work_order_id INTEGER NOT NULL,
            report_date TEXT,
            arrival_time TEXT,
            completion_time TEXT,
            issue_found TEXT,
            action_taken TEXT,
            materials_used TEXT,
            technician_notes TEXT,
            customer_notes TEXT,
            report_status TEXT DEFAULT 'draft',
            service_price REAL DEFAULT 0,
            technician_incentive REAL DEFAULT 0,
            region_allowance REAL DEFAULT 0,
            review_notes TEXT,
            submitted_by TEXT,
            submitted_at TEXT,
            reviewed_by TEXT,
            reviewed_at TEXT,
            status TEXT DEFAULT 'new',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    ensure_column(conn, "ops_work_orders", "request_date", "ALTER TABLE ops_work_orders ADD COLUMN request_date TEXT")
    ensure_column(conn, "ops_work_orders", "site_code", "ALTER TABLE ops_work_orders ADD COLUMN site_code TEXT")
    ensure_column(conn, "ops_work_orders", "site_name", "ALTER TABLE ops_work_orders ADD COLUMN site_name TEXT")
    ensure_column(conn, "ops_work_orders", "complaint_details", "ALTER TABLE ops_work_orders ADD COLUMN complaint_details TEXT")
    ensure_column(conn, "ops_work_orders", "required_materials", "ALTER TABLE ops_work_orders ADD COLUMN required_materials TEXT")
    ensure_column(conn, "ops_work_orders", "technician_id", "ALTER TABLE ops_work_orders ADD COLUMN technician_id INTEGER")
    ensure_column(conn, "ops_work_orders", "manager_id", "ALTER TABLE ops_work_orders ADD COLUMN manager_id INTEGER")
    ensure_column(conn, "ops_work_orders", "priority", "ALTER TABLE ops_work_orders ADD COLUMN priority TEXT DEFAULT 'normal'")
    ensure_column(conn, "ops_work_orders", "trip_required", "ALTER TABLE ops_work_orders ADD COLUMN trip_required INTEGER DEFAULT 1")
    ensure_column(conn, "ops_work_orders", "service_price", "ALTER TABLE ops_work_orders ADD COLUMN service_price REAL DEFAULT 0")
    ensure_column(conn, "ops_work_orders", "technician_incentive", "ALTER TABLE ops_work_orders ADD COLUMN technician_incentive REAL DEFAULT 0")
    ensure_column(conn, "ops_work_orders", "region_allowance", "ALTER TABLE ops_work_orders ADD COLUMN region_allowance REAL DEFAULT 0")
    ensure_column(conn, "ops_work_orders", "created_by", "ALTER TABLE ops_work_orders ADD COLUMN created_by TEXT")
    ensure_column(conn, "ops_work_orders", "updated_at", "ALTER TABLE ops_work_orders ADD COLUMN updated_at TEXT DEFAULT CURRENT_TIMESTAMP")
    ensure_column(conn, "ops_work_orders", "completed_at", "ALTER TABLE ops_work_orders ADD COLUMN completed_at TEXT")
    ensure_column(conn, "ops_work_orders", "closure_notes", "ALTER TABLE ops_work_orders ADD COLUMN closure_notes TEXT")
    ensure_column(conn, "ops_technician_reports", "report_date", "ALTER TABLE ops_technician_reports ADD COLUMN report_date TEXT")
    ensure_column(conn, "ops_technician_reports", "arrival_time", "ALTER TABLE ops_technician_reports ADD COLUMN arrival_time TEXT")
    ensure_column(conn, "ops_technician_reports", "completion_time", "ALTER TABLE ops_technician_reports ADD COLUMN completion_time TEXT")
    ensure_column(conn, "ops_technician_reports", "issue_found", "ALTER TABLE ops_technician_reports ADD COLUMN issue_found TEXT")
    ensure_column(conn, "ops_technician_reports", "action_taken", "ALTER TABLE ops_technician_reports ADD COLUMN action_taken TEXT")
    ensure_column(conn, "ops_technician_reports", "materials_used", "ALTER TABLE ops_technician_reports ADD COLUMN materials_used TEXT")
    ensure_column(conn, "ops_technician_reports", "technician_notes", "ALTER TABLE ops_technician_reports ADD COLUMN technician_notes TEXT")
    ensure_column(conn, "ops_technician_reports", "customer_notes", "ALTER TABLE ops_technician_reports ADD COLUMN customer_notes TEXT")
    ensure_column(conn, "ops_technician_reports", "report_status", "ALTER TABLE ops_technician_reports ADD COLUMN report_status TEXT DEFAULT 'draft'")
    ensure_column(conn, "ops_technician_reports", "service_price", "ALTER TABLE ops_technician_reports ADD COLUMN service_price REAL DEFAULT 0")
    ensure_column(conn, "ops_technician_reports", "technician_incentive", "ALTER TABLE ops_technician_reports ADD COLUMN technician_incentive REAL DEFAULT 0")
    ensure_column(conn, "ops_technician_reports", "region_allowance", "ALTER TABLE ops_technician_reports ADD COLUMN region_allowance REAL DEFAULT 0")
    ensure_column(conn, "ops_technician_reports", "review_notes", "ALTER TABLE ops_technician_reports ADD COLUMN review_notes TEXT")
    ensure_column(conn, "ops_technician_reports", "submitted_by", "ALTER TABLE ops_technician_reports ADD COLUMN submitted_by TEXT")
    ensure_column(conn, "ops_technician_reports", "submitted_at", "ALTER TABLE ops_technician_reports ADD COLUMN submitted_at TEXT")
    ensure_column(conn, "ops_technician_reports", "reviewed_by", "ALTER TABLE ops_technician_reports ADD COLUMN reviewed_by TEXT")
    ensure_column(conn, "ops_technician_reports", "reviewed_at", "ALTER TABLE ops_technician_reports ADD COLUMN reviewed_at TEXT")
    ensure_column(conn, "ops_action_materials", "action_id", "ALTER TABLE ops_action_materials ADD COLUMN action_id INTEGER")
    ensure_column(conn, "ops_action_materials", "line_no", "ALTER TABLE ops_action_materials ADD COLUMN line_no INTEGER DEFAULT 1")
    ensure_column(conn, "ops_action_materials", "item_id", "ALTER TABLE ops_action_materials ADD COLUMN item_id INTEGER")
    ensure_column(conn, "ops_action_materials", "item_code", "ALTER TABLE ops_action_materials ADD COLUMN item_code TEXT")
    ensure_column(conn, "ops_action_materials", "item_name", "ALTER TABLE ops_action_materials ADD COLUMN item_name TEXT")
    ensure_column(conn, "ops_action_materials", "uom", "ALTER TABLE ops_action_materials ADD COLUMN uom TEXT")
    ensure_column(conn, "ops_action_materials", "qty", "ALTER TABLE ops_action_materials ADD COLUMN qty REAL DEFAULT 0")
    ensure_column(conn, "ops_action_materials", "unit_cost", "ALTER TABLE ops_action_materials ADD COLUMN unit_cost REAL DEFAULT 0")
    ensure_column(conn, "ops_action_materials", "notes", "ALTER TABLE ops_action_materials ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_vehicle_rates", "ticket_open_price", "ALTER TABLE ops_vehicle_rates ADD COLUMN ticket_open_price REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_rates", "rental_supplier_id", "ALTER TABLE ops_vehicle_rates ADD COLUMN rental_supplier_id INTEGER")
    ensure_column(conn, "ops_vehicle_rates", "slab_name", "ALTER TABLE ops_vehicle_rates ADD COLUMN slab_name TEXT")
    ensure_column(conn, "ops_vehicle_rates", "second_slab_to_km", "ALTER TABLE ops_vehicle_rates ADD COLUMN second_slab_to_km REAL DEFAULT 300")
    ensure_column(conn, "ops_vehicle_rates", "km_rate_101_300", "ALTER TABLE ops_vehicle_rates ADD COLUMN km_rate_101_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_rates", "km_rate_over_300", "ALTER TABLE ops_vehicle_rates ADD COLUMN km_rate_over_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_rental_suppliers", "code", "ALTER TABLE ops_rental_suppliers ADD COLUMN code TEXT")
    ensure_column(conn, "ops_rental_suppliers", "contact_person", "ALTER TABLE ops_rental_suppliers ADD COLUMN contact_person TEXT")
    ensure_column(conn, "ops_rental_suppliers", "phone", "ALTER TABLE ops_rental_suppliers ADD COLUMN phone TEXT")
    ensure_column(conn, "ops_rental_suppliers", "notes", "ALTER TABLE ops_rental_suppliers ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_rental_suppliers", "is_active", "ALTER TABLE ops_rental_suppliers ADD COLUMN is_active INTEGER DEFAULT 1")
    ensure_column(conn, "ops_vehicles", "code", "ALTER TABLE ops_vehicles ADD COLUMN code TEXT")
    ensure_column(conn, "ops_vehicles", "vehicle_name", "ALTER TABLE ops_vehicles ADD COLUMN vehicle_name TEXT")
    ensure_column(conn, "ops_vehicles", "vehicle_type", "ALTER TABLE ops_vehicles ADD COLUMN vehicle_type TEXT")
    ensure_column(conn, "ops_vehicles", "rental_supplier_id", "ALTER TABLE ops_vehicles ADD COLUMN rental_supplier_id INTEGER")
    ensure_column(conn, "ops_vehicles", "vehicle_rate_id", "ALTER TABLE ops_vehicles ADD COLUMN vehicle_rate_id INTEGER")
    ensure_column(conn, "ops_vehicles", "pricing_slab_name", "ALTER TABLE ops_vehicles ADD COLUMN pricing_slab_name TEXT")
    ensure_column(conn, "ops_vehicles", "plate_no", "ALTER TABLE ops_vehicles ADD COLUMN plate_no TEXT")
    ensure_column(conn, "ops_vehicles", "driver_source", "ALTER TABLE ops_vehicles ADD COLUMN driver_source TEXT DEFAULT 'company_driver'")
    ensure_column(conn, "ops_vehicles", "supplier_driver_name", "ALTER TABLE ops_vehicles ADD COLUMN supplier_driver_name TEXT")
    ensure_column(conn, "ops_vehicles", "notes", "ALTER TABLE ops_vehicles ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_vehicles", "is_active", "ALTER TABLE ops_vehicles ADD COLUMN is_active INTEGER DEFAULT 1")
    ensure_column(conn, "ops_contracts", "contract_no", "ALTER TABLE ops_contracts ADD COLUMN contract_no TEXT")
    ensure_column(conn, "ops_contracts", "company_id", "ALTER TABLE ops_contracts ADD COLUMN company_id INTEGER")
    ensure_column(conn, "ops_contracts", "contract_name", "ALTER TABLE ops_contracts ADD COLUMN contract_name TEXT")
    ensure_column(conn, "ops_contracts", "pricing_method", "ALTER TABLE ops_contracts ADD COLUMN pricing_method TEXT DEFAULT 'standard'")
    ensure_column(conn, "ops_contracts", "start_date", "ALTER TABLE ops_contracts ADD COLUMN start_date TEXT")
    ensure_column(conn, "ops_contracts", "end_date", "ALTER TABLE ops_contracts ADD COLUMN end_date TEXT")
    ensure_column(conn, "ops_contracts", "status", "ALTER TABLE ops_contracts ADD COLUMN status TEXT DEFAULT 'active'")
    ensure_column(conn, "ops_contracts", "notes", "ALTER TABLE ops_contracts ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_action_price_versions", "contract_id", "ALTER TABLE ops_action_price_versions ADD COLUMN contract_id INTEGER")
    ensure_column(conn, "ops_action_price_versions", "action_id", "ALTER TABLE ops_action_price_versions ADD COLUMN action_id INTEGER")
    ensure_column(conn, "ops_action_price_versions", "version_name", "ALTER TABLE ops_action_price_versions ADD COLUMN version_name TEXT")
    ensure_column(conn, "ops_action_price_versions", "effective_from", "ALTER TABLE ops_action_price_versions ADD COLUMN effective_from TEXT")
    ensure_column(conn, "ops_action_price_versions", "effective_to", "ALTER TABLE ops_action_price_versions ADD COLUMN effective_to TEXT")
    ensure_column(conn, "ops_action_price_versions", "fuel_reference", "ALTER TABLE ops_action_price_versions ADD COLUMN fuel_reference TEXT")
    ensure_column(conn, "ops_action_price_versions", "action_price", "ALTER TABLE ops_action_price_versions ADD COLUMN action_price REAL DEFAULT 0")
    ensure_column(conn, "ops_action_price_versions", "technician_incentive", "ALTER TABLE ops_action_price_versions ADD COLUMN technician_incentive REAL DEFAULT 0")
    ensure_column(conn, "ops_action_price_versions", "region_allowance", "ALTER TABLE ops_action_price_versions ADD COLUMN region_allowance REAL DEFAULT 0")
    ensure_column(conn, "ops_action_price_versions", "is_active", "ALTER TABLE ops_action_price_versions ADD COLUMN is_active INTEGER DEFAULT 1")
    ensure_column(conn, "ops_action_price_versions", "notes", "ALTER TABLE ops_action_price_versions ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_vehicle_price_versions", "contract_id", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN contract_id INTEGER")
    ensure_column(conn, "ops_vehicle_price_versions", "vehicle_rate_id", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN vehicle_rate_id INTEGER")
    ensure_column(conn, "ops_vehicle_price_versions", "version_name", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN version_name TEXT")
    ensure_column(conn, "ops_vehicle_price_versions", "effective_from", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN effective_from TEXT")
    ensure_column(conn, "ops_vehicle_price_versions", "effective_to", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN effective_to TEXT")
    ensure_column(conn, "ops_vehicle_price_versions", "fuel_reference", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN fuel_reference TEXT")
    ensure_column(conn, "ops_vehicle_price_versions", "ticket_open_price", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN ticket_open_price REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_price_versions", "second_slab_to_km", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN second_slab_to_km REAL DEFAULT 300")
    ensure_column(conn, "ops_vehicle_price_versions", "km_rate_101_300", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN km_rate_101_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_price_versions", "km_rate_over_300", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN km_rate_over_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_price_versions", "waiting_hour_rate", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN waiting_hour_rate REAL DEFAULT 0")
    ensure_column(conn, "ops_vehicle_price_versions", "is_active", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN is_active INTEGER DEFAULT 1")
    ensure_column(conn, "ops_vehicle_price_versions", "notes", "ALTER TABLE ops_vehicle_price_versions ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_tickets", "ticket_no", "ALTER TABLE ops_tickets ADD COLUMN ticket_no TEXT")
    ensure_column(conn, "ops_tickets", "company_id", "ALTER TABLE ops_tickets ADD COLUMN company_id INTEGER")
    ensure_column(conn, "ops_tickets", "contract_id", "ALTER TABLE ops_tickets ADD COLUMN contract_id INTEGER")
    ensure_column(conn, "ops_tickets", "ticket_date", "ALTER TABLE ops_tickets ADD COLUMN ticket_date TEXT")
    ensure_column(conn, "ops_tickets", "fault_type_id", "ALTER TABLE ops_tickets ADD COLUMN fault_type_id INTEGER")
    ensure_column(conn, "ops_tickets", "site_code", "ALTER TABLE ops_tickets ADD COLUMN site_code TEXT")
    ensure_column(conn, "ops_tickets", "site_name", "ALTER TABLE ops_tickets ADD COLUMN site_name TEXT")
    ensure_column(conn, "ops_tickets", "priority", "ALTER TABLE ops_tickets ADD COLUMN priority TEXT DEFAULT 'normal'")
    ensure_column(conn, "ops_tickets", "request_channel", "ALTER TABLE ops_tickets ADD COLUMN request_channel TEXT")
    ensure_column(conn, "ops_tickets", "complaint_details", "ALTER TABLE ops_tickets ADD COLUMN complaint_details TEXT")
    ensure_column(conn, "ops_tickets", "status", "ALTER TABLE ops_tickets ADD COLUMN status TEXT DEFAULT 'open'")
    ensure_column(conn, "ops_tickets", "created_by", "ALTER TABLE ops_tickets ADD COLUMN created_by TEXT")
    ensure_column(conn, "ops_trip_tickets", "trip_no", "ALTER TABLE ops_trip_tickets ADD COLUMN trip_no TEXT")
    ensure_column(conn, "ops_trip_tickets", "trip_date", "ALTER TABLE ops_trip_tickets ADD COLUMN trip_date TEXT")
    ensure_column(conn, "ops_trip_tickets", "vehicle_id", "ALTER TABLE ops_trip_tickets ADD COLUMN vehicle_id INTEGER")
    ensure_column(conn, "ops_trip_tickets", "rental_supplier_id", "ALTER TABLE ops_trip_tickets ADD COLUMN rental_supplier_id INTEGER")
    ensure_column(conn, "ops_trip_tickets", "vehicle_type", "ALTER TABLE ops_trip_tickets ADD COLUMN vehicle_type TEXT")
    ensure_column(conn, "ops_trip_tickets", "vehicle_rate_id", "ALTER TABLE ops_trip_tickets ADD COLUMN vehicle_rate_id INTEGER")
    ensure_column(conn, "ops_trip_tickets", "driver_source", "ALTER TABLE ops_trip_tickets ADD COLUMN driver_source TEXT DEFAULT 'company_driver'")
    ensure_column(conn, "ops_trip_tickets", "driver_employee_id", "ALTER TABLE ops_trip_tickets ADD COLUMN driver_employee_id INTEGER")
    ensure_column(conn, "ops_trip_tickets", "supplier_driver_name", "ALTER TABLE ops_trip_tickets ADD COLUMN supplier_driver_name TEXT")
    ensure_column(conn, "ops_trip_tickets", "status", "ALTER TABLE ops_trip_tickets ADD COLUMN status TEXT DEFAULT 'draft'")
    ensure_column(conn, "ops_trip_tickets", "start_odometer", "ALTER TABLE ops_trip_tickets ADD COLUMN start_odometer REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "start_photo_path", "ALTER TABLE ops_trip_tickets ADD COLUMN start_photo_path TEXT")
    ensure_column(conn, "ops_trip_tickets", "end_odometer", "ALTER TABLE ops_trip_tickets ADD COLUMN end_odometer REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "end_photo_path", "ALTER TABLE ops_trip_tickets ADD COLUMN end_photo_path TEXT")
    ensure_column(conn, "ops_trip_tickets", "waiting_hours", "ALTER TABLE ops_trip_tickets ADD COLUMN waiting_hours REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "total_km", "ALTER TABLE ops_trip_tickets ADD COLUMN total_km REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "ticket_open_price", "ALTER TABLE ops_trip_tickets ADD COLUMN ticket_open_price REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "second_slab_to_km", "ALTER TABLE ops_trip_tickets ADD COLUMN second_slab_to_km REAL DEFAULT 300")
    ensure_column(conn, "ops_trip_tickets", "km_rate_101_300", "ALTER TABLE ops_trip_tickets ADD COLUMN km_rate_101_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "km_rate_over_300", "ALTER TABLE ops_trip_tickets ADD COLUMN km_rate_over_300 REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "waiting_hour_rate", "ALTER TABLE ops_trip_tickets ADD COLUMN waiting_hour_rate REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "total_cost", "ALTER TABLE ops_trip_tickets ADD COLUMN total_cost REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "allocated_cost_per_work_order", "ALTER TABLE ops_trip_tickets ADD COLUMN allocated_cost_per_work_order REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "driver_commission_pct", "ALTER TABLE ops_trip_tickets ADD COLUMN driver_commission_pct REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "driver_commission_amount", "ALTER TABLE ops_trip_tickets ADD COLUMN driver_commission_amount REAL DEFAULT 0")
    ensure_column(conn, "ops_trip_tickets", "notes", "ALTER TABLE ops_trip_tickets ADD COLUMN notes TEXT")
    ensure_column(conn, "ops_trip_tickets", "movement_notes", "ALTER TABLE ops_trip_tickets ADD COLUMN movement_notes TEXT")
    ensure_column(conn, "ops_trip_tickets", "accounting_notes", "ALTER TABLE ops_trip_tickets ADD COLUMN accounting_notes TEXT")
    ensure_column(conn, "ops_trip_tickets", "created_by", "ALTER TABLE ops_trip_tickets ADD COLUMN created_by TEXT")
    ensure_column(conn, "ops_trip_tickets", "movement_closed_by", "ALTER TABLE ops_trip_tickets ADD COLUMN movement_closed_by TEXT")
    ensure_column(conn, "ops_trip_tickets", "approved_by", "ALTER TABLE ops_trip_tickets ADD COLUMN approved_by TEXT")
    ensure_column(conn, "ops_trip_tickets", "completed_at", "ALTER TABLE ops_trip_tickets ADD COLUMN completed_at TEXT")
    ensure_column(conn, "ops_trip_tickets", "approved_at", "ALTER TABLE ops_trip_tickets ADD COLUMN approved_at TEXT")
    ensure_column(conn, "ops_trip_work_orders", "trip_id", "ALTER TABLE ops_trip_work_orders ADD COLUMN trip_id INTEGER")
    ensure_column(conn, "ops_trip_work_orders", "work_order_id", "ALTER TABLE ops_trip_work_orders ADD COLUMN work_order_id INTEGER")
    ensure_column(conn, "ops_trip_work_orders", "line_no", "ALTER TABLE ops_trip_work_orders ADD COLUMN line_no INTEGER DEFAULT 1")
    ensure_column(conn, "ops_trip_work_orders", "allocated_cost", "ALTER TABLE ops_trip_work_orders ADD COLUMN allocated_cost REAL DEFAULT 0")
    ensure_column(conn, "employees", "is_driver", "ALTER TABLE employees ADD COLUMN is_driver INTEGER DEFAULT 0")
    ensure_column(conn, "employees", "trip_commission_pct", "ALTER TABLE employees ADD COLUMN trip_commission_pct REAL DEFAULT 0")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_ops_report_work_order ON ops_technician_reports(work_order_id)")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_ops_trip_work_order_unique ON ops_trip_work_orders(trip_id, work_order_id)")
    conn.commit()
    conn.close()


ensure_tables()


def next_code(table_name: str, prefix: str):
    conn = get_conn()
    row = conn.execute(
        f"SELECT code FROM {table_name} WHERE COALESCE(code, '') <> '' ORDER BY id DESC LIMIT 1"
    ).fetchone()
    conn.close()
    if not row or not row["code"]:
        return f"{prefix}-0001"
    last = safe(row["code"])
    try:
        num = int(last.split("-")[-1])
    except Exception:
        num = 0
    return f"{prefix}-{num + 1:04d}"


def next_number(table_name: str, field_name: str, prefix: str):
    conn = get_conn()
    row = conn.execute(
        f"SELECT {field_name} AS serial_value FROM {table_name} WHERE COALESCE({field_name}, '') <> '' ORDER BY id DESC LIMIT 1"
    ).fetchone()
    conn.close()
    if not row or not row["serial_value"]:
        return f"{prefix}-0001"
    last = safe(row["serial_value"])
    try:
        num = int(last.split("-")[-1])
    except Exception:
        num = 0
    return f"{prefix}-{num + 1:04d}"


def count_rows(table_name: str):
    conn = get_conn()
    row = conn.execute(f"SELECT COUNT(*) AS cnt FROM {table_name}").fetchone()
    conn.close()
    return int(row["cnt"] or 0)


def current_form_notice(notice: str):
    if safe(notice).lower() == "saved":
        return '<div class="msg success">Saved successfully.</div>'
    if safe(notice).lower() == "updated":
        return '<div class="msg success">Updated successfully.</div>'
    if safe(notice):
        return f'<div class="msg success">{safe(notice)}</div>'
    return ""


def status_options(selected=""):
    html = ""
    for value, label in WORK_ORDER_STATUSES:
        sel = "selected" if safe(selected) == value else ""
        html += f"<option value='{value}' {sel}>{label}</option>"
    return html


def status_chip(status: str):
    state = safe(status).lower()
    color = "blue"
    if state in ["approved", "closed"]:
        color = "green"
    elif state in ["rejected"]:
        color = "red"
    elif state in ["submitted", "in_progress"]:
        color = "orange"
    elif state in ["assigned"]:
        color = "blue"
    return f"<span class='status-chip {color}'>{safe(status).replace('_', ' ').title()}</span>"


def parse_csv_rows(file_bytes):
    text = file_bytes.decode("utf-8-sig", errors="ignore")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(r) for r in reader]


def parse_xlsx_rows(file_bytes):
    if load_workbook is None:
        raise Exception("Excel import is not available right now.")
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [safe(h) for h in rows[0]]
    result = []
    for data_row in rows[1:]:
        item = {}
        for i, header in enumerate(headers):
            if not header:
                continue
            item[header] = "" if i >= len(data_row) or data_row[i] is None else str(data_row[i])
        result.append(item)
    return result


def next_action_code():
    return next_code("ops_service_catalog", "ACT")


def next_ticket_no():
    return next_number("ops_tickets", "ticket_no", "TKT")


def next_contract_no():
    return next_number("ops_contracts", "contract_no", "CTR")


def next_trip_no():
    return next_number("ops_trip_tickets", "trip_no", "TRP")


def next_rental_supplier_code():
    return next_code("ops_rental_suppliers", "SUP")


def next_vehicle_code():
    return next_code("ops_vehicles", "CAR")


def normalize_action_import_row(row):
    return {
        "code": safe(row.get("code") or row.get("action_code") or row.get("item_code")),
        "name": safe(row.get("name") or row.get("action_name") or row.get("action")),
        "service_category": safe(row.get("service_category") or row.get("action_category") or "field_maintenance"),
        "unit_price": safe(row.get("unit_price") or row.get("price") or "0"),
        "technician_incentive": safe(row.get("technician_incentive") or row.get("incentive") or "0"),
        "default_region_level": safe(row.get("default_region_level") or row.get("region_level") or "zone_1"),
        "default_duration_hours": safe(row.get("default_duration_hours") or row.get("duration_hours") or "0"),
        "notes": safe(row.get("notes")),
        "is_active": safe(row.get("is_active") or "1"),
    }


def get_item_rows():
    conn = get_conn()
    rows = conn.execute("""
        SELECT id, code, name, uom
        FROM items
        WHERE COALESCE(is_active, 1) = 1
          AND LOWER(COALESCE(item_type, 'stock_item')) = 'stock_item'
        ORDER BY code, name
    """).fetchall()
    conn.close()
    return rows


def company_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT id, code, name
        FROM ops_contract_companies
        WHERE COALESCE(is_active, 1) = 1
        ORDER BY name
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Company --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['code'])} - {safe(row['name'])}" if safe(row["code"]) else safe(row["name"])
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def contract_options(selected_id=0, company_id=0):
    conn = get_conn()
    sql = """
        SELECT c.id, c.contract_no, c.contract_name, co.name AS company_name
        FROM ops_contracts c
        LEFT JOIN ops_contract_companies co ON co.id = c.company_id
        WHERE LOWER(COALESCE(c.status, 'active')) = 'active'
    """
    params = []
    if safe_int(company_id) > 0:
        sql += " AND c.company_id = ?"
        params.append(safe_int(company_id))
    sql += " ORDER BY c.id DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    html = "<option value=''>-- Select Contract --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['contract_no'])} - {safe(row['contract_name'])}"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def action_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT id, code, name
        FROM ops_service_catalog
        WHERE COALESCE(is_active, 1) = 1
        ORDER BY name
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Action --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['code'])} - {safe(row['name'])}"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def fault_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT id, code, name
        FROM ops_fault_types
        WHERE COALESCE(is_active, 1) = 1
        ORDER BY name
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Fault --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['code'])} - {safe(row['name'])}"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def item_options(selected_id=0):
    html = "<option value=''>-- Select Item --</option>"
    for row in get_item_rows():
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['code'])} - {safe(row['name'])} ({safe(row['uom'])})"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def employee_display_name(row):
    return safe(
        row["employee_name"]
        if "employee_name" in row.keys() and row["employee_name"]
        else row["name"]
        if "name" in row.keys() and row["name"]
        else row["full_name"]
        if "full_name" in row.keys() and row["full_name"]
        else row["code"]
        if "code" in row.keys() and row["code"]
        else ""
    )


def employee_options(selected_id=0, driver_only=False):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            id,
            code,
            name,
            employee_name,
            full_name,
            department,
            job_title,
            COALESCE(is_driver, 0) AS is_driver
        FROM employees
        WHERE COALESCE(is_active, 1) = 1
        ORDER BY COALESCE(employee_name, name, full_name, code)
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Employee --</option>"
    filtered_rows = []
    for row in rows:
        if not driver_only:
            filtered_rows.append(row)
            continue
        job_title = safe(row["job_title"]).lower()
        department = safe(row["department"]).lower()
        if int(row["is_driver"] or 0) == 1 or "driver" in job_title or "transport" in department or "حركة" in department:
            filtered_rows.append(row)
    if driver_only and not filtered_rows:
        filtered_rows = rows
    for row in filtered_rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = employee_display_name(row)
        if safe(row["code"]):
            label = f"{safe(row['code'])} - {label}"
        if safe(row["job_title"]):
            label = f"{label} ({safe(row['job_title'])})"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def rental_supplier_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT id, code, name
        FROM ops_rental_suppliers
        WHERE COALESCE(is_active, 1) = 1
        ORDER BY name
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Rental Office --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = f"{safe(row['code'])} - {safe(row['name'])}" if safe(row["code"]) else safe(row["name"])
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def vehicle_rate_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            vr.id,
            vr.code,
            vr.vehicle_type,
            vr.slab_name,
            rs.name AS supplier_name
        FROM ops_vehicle_rates vr
        LEFT JOIN ops_rental_suppliers rs ON rs.id = vr.rental_supplier_id
        WHERE COALESCE(vr.is_active, 1) = 1
        ORDER BY COALESCE(rs.name, ''), vr.vehicle_type, vr.slab_name, vr.code
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Vehicle Rate --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = safe(row["code"])
        extras = [x for x in [safe(row["supplier_name"]), safe(row["vehicle_type"]), safe(row["slab_name"])] if x]
        if extras:
            label = f"{label} - {' / '.join(extras)}"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def vehicle_options(selected_id=0):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            v.id,
            v.code,
            v.vehicle_name,
            v.vehicle_type,
            v.plate_no,
            rs.name AS supplier_name
        FROM ops_vehicles v
        LEFT JOIN ops_rental_suppliers rs ON rs.id = v.rental_supplier_id
        WHERE COALESCE(v.is_active, 1) = 1
        ORDER BY v.code, v.vehicle_name, v.vehicle_type
    """).fetchall()
    conn.close()
    html = "<option value=''>-- Select Vehicle --</option>"
    for row in rows:
        sel = "selected" if str(selected_id or "") == str(row["id"]) else ""
        label = safe(row["code"]) or safe(row["vehicle_name"]) or safe(row["vehicle_type"])
        extras = []
        if safe(row["vehicle_type"]):
            extras.append(safe(row["vehicle_type"]))
        if safe(row["supplier_name"]):
            extras.append(safe(row["supplier_name"]))
        if safe(row["plate_no"]):
            extras.append(safe(row["plate_no"]))
        if extras:
            label = f"{label} ({' / '.join(extras)})"
        html += f"<option value='{row['id']}' {sel}>{label}</option>"
    return html


def driver_source_options(selected="company_driver"):
    html = ""
    for value, label in DRIVER_SOURCE_OPTIONS:
        sel = "selected" if safe(selected) == value else ""
        html += f"<option value='{value}' {sel}>{label}</option>"
    return html


def trip_status_chip(status: str):
    state = safe(status).lower()
    color = "gray"
    if state == "approved":
        color = "green"
    elif state == "completed":
        color = "orange"
    elif state == "dispatched":
        color = "blue"
    return f"<span class='status-chip {color}'>{safe(status).replace('_', ' ').title()}</span>"


def trip_upload_dir():
    return os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "static", "uploads", "trip-meters"))


async def save_trip_photo(file: UploadFile, prefix: str):
    filename = safe(file.filename)
    if not filename:
        return ""
    ext = os.path.splitext(filename)[1] or ".jpg"
    folder = trip_upload_dir()
    os.makedirs(folder, exist_ok=True)
    saved_name = f"{prefix}_{uuid.uuid4().hex[:12]}{ext}"
    target = os.path.join(folder, saved_name)
    with open(target, "wb") as output:
        output.write(await file.read())
    return f"/static/uploads/trip-meters/{saved_name}"


def get_vehicle_row(vehicle_id: int):
    conn = get_conn()
    row = conn.execute("""
        SELECT
            v.*,
            rs.name AS supplier_name,
            vr.id AS vehicle_rate_id_resolved,
            vr.code AS vehicle_rate_code,
            vr.slab_name AS rate_slab_name,
            vr.ticket_open_price,
            vr.second_slab_to_km,
            vr.km_rate_101_300,
            vr.km_rate_over_300,
            vr.waiting_hour_rate
        FROM ops_vehicles v
        LEFT JOIN ops_rental_suppliers rs ON rs.id = v.rental_supplier_id
        LEFT JOIN ops_vehicle_rates vr ON vr.id = (
            SELECT vr2.id
            FROM ops_vehicle_rates vr2
            WHERE COALESCE(vr2.is_active, 1) = 1
              AND COALESCE(vr2.rental_supplier_id, 0) = COALESCE(v.rental_supplier_id, 0)
              AND LOWER(COALESCE(vr2.vehicle_type, '')) = LOWER(COALESCE(v.vehicle_type, ''))
              AND LOWER(COALESCE(vr2.slab_name, 'standard')) = LOWER(COALESCE(CASE WHEN COALESCE(v.pricing_slab_name, '') = '' THEN 'standard' ELSE v.pricing_slab_name END, 'standard'))
            ORDER BY vr2.id DESC
            LIMIT 1
        )
        WHERE v.id = ?
        LIMIT 1
    """, (vehicle_id,)).fetchone()
    if row and not row["vehicle_rate_id_resolved"] and safe_int(row["vehicle_rate_id"]) > 0:
        row = conn.execute("""
            SELECT
                v.*,
                rs.name AS supplier_name,
                vr.id AS vehicle_rate_id_resolved,
                vr.code AS vehicle_rate_code,
                vr.slab_name AS rate_slab_name,
                vr.ticket_open_price,
                vr.second_slab_to_km,
                vr.km_rate_101_300,
                vr.km_rate_over_300,
                vr.waiting_hour_rate
            FROM ops_vehicles v
            LEFT JOIN ops_rental_suppliers rs ON rs.id = v.rental_supplier_id
            LEFT JOIN ops_vehicle_rates vr ON vr.id = v.vehicle_rate_id
            WHERE v.id = ?
            LIMIT 1
        """, (vehicle_id,)).fetchone()
    conn.close()
    return row


def get_driver_row(employee_id: int):
    conn = get_conn()
    row = conn.execute("""
        SELECT
            id,
            code,
            name,
            employee_name,
            full_name,
            job_title,
            COALESCE(trip_commission_pct, 0) AS trip_commission_pct
        FROM employees
        WHERE id = ?
        LIMIT 1
    """, (employee_id,)).fetchone()
    conn.close()
    return row


def work_order_choices(selected_ids=None):
    selected_ids = selected_ids or []
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            wo.id,
            wo.work_order_no,
            wo.site_name,
            wo.status,
            wo.request_date,
            c.name AS company_name
        FROM ops_work_orders wo
        LEFT JOIN ops_contract_companies c ON c.id = wo.company_id
        WHERE COALESCE(wo.trip_required, 1) = 1
        ORDER BY wo.id DESC
    """).fetchall()
    conn.close()
    html = ""
    for row in rows:
        checked = "checked" if str(row["id"]) in [str(x) for x in selected_ids] else ""
        label = safe(row["work_order_no"]) or f"WO-{row['id']}"
        meta = " / ".join([x for x in [safe(row["company_name"]), safe(row["site_name"]), safe(row["status"]).title()] if x])
        html += f"""
        <label class="checkbox-card">
            <input type="checkbox" name="work_order_ids" value="{row['id']}" {checked}>
            <span><b>{label}</b><br><small>{meta or 'No extra details'}</small></span>
        </label>
        """
    if not html:
        html = "<div class='empty-note'>No work orders available yet. Create work orders first, then attach them to the trip.</div>"
    return html


def calculate_trip_cost(total_km, waiting_hours, ticket_open_price, second_slab_to_km, km_rate_101_300, km_rate_over_300, waiting_hour_rate):
    km_value = q2(total_km)
    waiting_value = q2(waiting_hours)
    waiting_cost = waiting_value * q2(waiting_hour_rate)
    second_slab_limit = q2(second_slab_to_km)
    if second_slab_limit < Decimal("100.00"):
        second_slab_limit = Decimal("300.00")
    if km_value <= Decimal("100.00"):
        total = q2(ticket_open_price) + waiting_cost
    elif km_value <= second_slab_limit:
        total = q2(ticket_open_price) + ((km_value - Decimal("100.00")) * q2(km_rate_101_300)) + waiting_cost
    else:
        total = (km_value * q2(km_rate_over_300)) + waiting_cost
    return q2(total)


def get_item_by_id(item_id: int):
    conn = get_conn()
    row = conn.execute("""
        SELECT id, code, name, uom
        FROM items
        WHERE id = ?
        LIMIT 1
    """, (item_id,)).fetchone()
    conn.close()
    return row


def get_item_last_cost(item_id: int):
    conn = get_conn()
    row = conn.execute("""
        SELECT unit_cost
        FROM stock_ledger
        WHERE item_id = ?
          AND COALESCE(unit_cost, 0) > 0
        ORDER BY id DESC
        LIMIT 1
    """, (item_id,)).fetchone()
    conn.close()
    return q2(row["unit_cost"] if row else 0)


def get_action_materials(conn, action_id: int):
    return conn.execute("""
        SELECT *
        FROM ops_action_materials
        WHERE action_id = ?
        ORDER BY line_no, id
    """, (action_id,)).fetchall()


def service_category_options(selected=""):
    html = ""
    for value, label in SERVICE_CATEGORIES:
        sel = "selected" if safe(selected) == value else ""
        html += f"<option value='{value}' {sel}>{label}</option>"
    return html


def zone_level_options(selected=""):
    html = ""
    for value, label in ZONE_LEVELS:
        sel = "selected" if safe(selected) == value else ""
        html += f"<option value='{value}' {sel}>{label}</option>"
    return html


def active_checkbox(checked=True):
    return "checked" if checked else ""


def module_cards(cards):
    html = '<div class="card-grid">'
    for title, href, icon, desc in cards:
        icon_html = f'<img src="{icon}" alt="{title} icon">' if icon.startswith("/static/") else icon
        html += f"""
        <a href="{href}" class="module-card">
            <div class="module-card-icon">{icon_html}</div>
            <div class="module-card-title">{title}</div>
            <div class="module-card-sub">{desc}</div>
        </a>
        """
    html += "</div>"
    return html


def render_ops_page(request: Request, title: str, content: str):
    return HTMLResponse(render_page(title, content, current_path=str(request.url.path)))


@router.get("/ui/projects")
def projects_legacy_redirect():
    return RedirectResponse("/ui/operations", status_code=302)


@router.get("/ui/operations", response_class=HTMLResponse)
def operations_root(request: Request):
    setup_cards = [
        ("Tickets", "/ui/operations/tickets", "/static/icons/journal.svg", "Open customer operational tickets before work order assignment."),
        ("Contracts", "/ui/operations/contracts", "/static/icons/customers.svg", "Different commercial agreements and pricing methods by company."),
        ("Pricing Versions", "/ui/operations/pricing-versions", "/static/icons/reports.svg", "Versioned action and vehicle prices by contract and effective date."),
        ("Trip Tickets", "/ui/operations/trips", "/static/icons/goods-receipts.svg", "Draft, movement, and accounting approval for vehicle trips linked to work orders."),
        ("Vehicles", "/ui/operations/vehicles", "/static/icons/goods-receipts.svg", "Vehicle master with code, rental office, driver source, and linked rate."),
        ("Rental Offices", "/ui/operations/rental-suppliers", "/static/icons/vendors.svg", "Rental suppliers or car offices that provide vehicles by location."),
        ("Work Orders", "/ui/operations/work-orders", "/static/icons/journal.svg", "Create, assign, track, and review field maintenance and workshop jobs."),
        ("Contract Companies", "/ui/operations/companies", "/static/icons/customers.svg", "Companies that send faults, modules, assemblies, and customer custody stock."),
        ("Fault Types", "/ui/operations/fault-types", "/static/icons/reports.svg", "Known fault codes used when opening tickets and classifying field issues."),
        ("Regions", "/ui/operations/regions", "/static/icons/reports.svg", "Zone allowances for field visits and technician bonus by area."),
        ("Action Catalog", "/ui/operations/action-catalog", "/static/icons/customer-invoices.svg", "Priced actions with import template and fixed raw materials for work order costing."),
        ("Vehicle Rates", "/ui/operations/vehicle-rates", "/static/icons/goods-receipts.svg", "Kilometer and waiting-hour rates for trip tickets."),
    ]
    summary = f"""
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Operations</h2>
                <div class="section-note">Field maintenance, workshop repairs, cabinet assembly, and customer custody stock in one workflow.</div>
            </div>
            <div class="table-summary">
                <span class="summary-pill">Companies: {count_rows('ops_contract_companies')}</span>
                <span class="summary-pill">Contracts: {count_rows('ops_contracts')}</span>
                <span class="summary-pill">Tickets: {count_rows('ops_tickets')}</span>
                <span class="summary-pill">Fault Types: {count_rows('ops_fault_types')}</span>
                <span class="summary-pill">Regions: {count_rows('ops_regions')}</span>
                <span class="summary-pill">Actions: {count_rows('ops_service_catalog')}</span>
                <span class="summary-pill">Rental Offices: {count_rows('ops_rental_suppliers')}</span>
                <span class="summary-pill">Vehicles: {count_rows('ops_vehicles')}</span>
                <span class="summary-pill">Vehicle Rates: {count_rows('ops_vehicle_rates')}</span>
                <span class="summary-pill">Trips: {count_rows('ops_trip_tickets')}</span>
                <span class="summary-pill">Work Orders: {count_rows('ops_work_orders')}</span>
            </div>
        </div>
    </div>
    """
    workflow = """
    <div class="card">
        <h3 class="sub-title">Operational Flow</h3>
        <div class="section-note">Trip workflow is now live first: draft by technical manager, completion by movement manager, and final approval by accounting before transport cost is distributed.</div>
        <div class="form-grid">
            <div class="form-group"><label>Field Maintenance</label><input value="Ticket -> Work Order -> Trip Ticket -> Technician Report -> Manager Review -> Invoice" readonly></div>
            <div class="form-group"><label>Workshop Repairs</label><input value="Customer Intake -> Repair Order -> Parts Usage -> Test -> Return to Customer" readonly></div>
            <div class="form-group"><label>Cabinet Assembly</label><input value="Assembly Request -> Components -> Build -> Completion -> Delivery" readonly></div>
            <div class="form-group"><label>Customer Custody</label><input value="Custody Receipt -> Issue to Work Order -> Balance Statement by Customer" readonly></div>
        </div>
    </div>
    """
    content = summary + '<div class="card"><h3 class="sub-title">Master Data</h3>' + module_cards(setup_cards) + "</div>" + workflow
    return render_ops_page(request, "Operations", content)


@router.get("/ui/operations/companies", response_class=HTMLResponse)
def companies_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ops_contract_companies ORDER BY id DESC").fetchall()
    edit_row = conn.execute("SELECT * FROM ops_contract_companies WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_code("ops_contract_companies", "COMP"),
        "name": "",
        "contact_person": "",
        "phone": "",
        "email": "",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        status = "Active" if int(row["is_active"] or 0) == 1 else "Inactive"
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['name'])}</td>
            <td>{safe(row['contact_person'])}</td>
            <td>{safe(row['phone'])}</td>
            <td>{safe(row['email'])}</td>
            <td>{status}</td>
            <td><a class="btn blue" href="/ui/operations/companies?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='7' style='text-align:center;'>No contract companies added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Contract Companies</h2>
                <div class="section-note">Companies that open faults, send modules for repair, request cabinet assembly, or hand over custody stock.</div>
            </div>
            <a class="btn gray" href="/ui/operations">Back to Operations</a>
        </div>
        <form method="post" action="/ui/operations/companies/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Company Name</label><input name="name" value="{safe(form_values.get('name'))}" required></div>
                <div class="form-group"><label>Contact Person</label><input name="contact_person" value="{safe(form_values.get('contact_person'))}"></div>
                <div class="form-group"><label>Phone</label><input name="phone" value="{safe(form_values.get('phone'))}"></div>
                <div class="form-group"><label>Email</label><input name="email" value="{safe(form_values.get('email'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Company</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Company" if safe_int(form_values.get('id')) > 0 else "Save Company"}</button>
                <a class="btn gray" href="/ui/operations/companies">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Company List</h3>
        <table>
            <tr><th>Code</th><th>Name</th><th>Contact</th><th>Phone</th><th>Email</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Contract Companies", html)


@router.post("/ui/operations/companies/save")
def save_company(
    row_id: int = Form(0),
    code: str = Form(""),
    name: str = Form(""),
    contact_person: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    if row_id > 0:
        conn.execute("""
            UPDATE ops_contract_companies
            SET code = ?, name = ?, contact_person = ?, phone = ?, email = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (safe(code), safe(name), safe(contact_person), safe(phone), safe(email), safe(notes), active_flag, row_id))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_contract_companies (code, name, contact_person, phone, email, notes, is_active)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (safe(code) or next_code("ops_contract_companies", "COMP"), safe(name), safe(contact_person), safe(phone), safe(email), safe(notes), active_flag))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/companies?notice={notice}", status_code=303)


@router.get("/ui/operations/fault-types", response_class=HTMLResponse)
def fault_types_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ops_fault_types ORDER BY id DESC").fetchall()
    edit_row = conn.execute("SELECT * FROM ops_fault_types WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_code("ops_fault_types", "FLT"),
        "name": "",
        "service_category": "field_maintenance",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['name'])}</td>
            <td>{safe(row['service_category']).replace('_', ' ').title()}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td><a class="btn blue" href="/ui/operations/fault-types?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='5' style='text-align:center;'>No fault types added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Fault Types</h2>
                <div class="section-note">Known fault names used when opening tickets. Pricing and incentive should be configured on the Action side.</div>
            </div>
            <a class="btn gray" href="/ui/operations">Back to Operations</a>
        </div>
        <form method="post" action="/ui/operations/fault-types/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Fault Name</label><input name="name" value="{safe(form_values.get('name'))}" required></div>
                <div class="form-group"><label>Service Category</label><select name="service_category">{service_category_options(form_values.get('service_category'))}</select></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Fault</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Fault Type" if safe_int(form_values.get('id')) > 0 else "Save Fault Type"}</button>
                <a class="btn gray" href="/ui/operations/fault-types">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Fault Type List</h3>
        <table>
            <tr><th>Code</th><th>Fault</th><th>Category</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Fault Types", html)


@router.post("/ui/operations/fault-types/save")
def save_fault_type(
    row_id: int = Form(0),
    code: str = Form(""),
    name: str = Form(""),
    service_category: str = Form("field_maintenance"),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    existing = conn.execute("SELECT default_service_price, default_incentive FROM ops_fault_types WHERE id = ? LIMIT 1", (row_id,)).fetchone() if row_id > 0 else None
    if row_id > 0:
        conn.execute("""
            UPDATE ops_fault_types
            SET code = ?, name = ?, service_category = ?, default_service_price = ?, default_incentive = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (
            safe(code),
            safe(name),
            safe(service_category),
            float(q2(existing["default_service_price"] if existing else 0)),
            float(q2(existing["default_incentive"] if existing else 0)),
            safe(notes),
            active_flag,
            row_id,
        ))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_fault_types (code, name, service_category, default_service_price, default_incentive, notes, is_active)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            safe(code) or next_code("ops_fault_types", "FLT"),
            safe(name),
            safe(service_category),
            0.0,
            0.0,
            safe(notes),
            active_flag,
        ))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/fault-types?notice={notice}", status_code=303)


@router.get("/ui/operations/regions", response_class=HTMLResponse)
def regions_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ops_regions ORDER BY id DESC").fetchall()
    edit_row = conn.execute("SELECT * FROM ops_regions WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_code("ops_regions", "REG"),
        "name": "",
        "zone_level": "zone_1",
        "allowance_amount": "0.00",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['name'])}</td>
            <td>{safe(row['zone_level']).replace('_', ' ').title()}</td>
            <td>{money(row['allowance_amount'])}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td><a class="btn blue" href="/ui/operations/regions?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='6' style='text-align:center;'>No regions added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Regions</h2>
                <div class="section-note">Area bands used for technician allowance and trip costing by destination.</div>
            </div>
            <a class="btn gray" href="/ui/operations">Back to Operations</a>
        </div>
        <form method="post" action="/ui/operations/regions/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Region Name</label><input name="name" value="{safe(form_values.get('name'))}" required></div>
                <div class="form-group"><label>Zone Level</label><select name="zone_level">{zone_level_options(form_values.get('zone_level'))}</select></div>
                <div class="form-group"><label>Allowance Amount</label><input name="allowance_amount" value="{safe(form_values.get('allowance_amount'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Region</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Region" if safe_int(form_values.get('id')) > 0 else "Save Region"}</button>
                <a class="btn gray" href="/ui/operations/regions">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Region List</h3>
        <table>
            <tr><th>Code</th><th>Region</th><th>Zone Level</th><th>Allowance</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Regions", html)


@router.post("/ui/operations/regions/save")
def save_region(
    row_id: int = Form(0),
    code: str = Form(""),
    name: str = Form(""),
    zone_level: str = Form("zone_1"),
    allowance_amount: str = Form("0"),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    if row_id > 0:
        conn.execute("""
            UPDATE ops_regions
            SET code = ?, name = ?, zone_level = ?, allowance_amount = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (safe(code), safe(name), safe(zone_level), float(q2(allowance_amount)), safe(notes), active_flag, row_id))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_regions (code, name, zone_level, allowance_amount, notes, is_active)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (safe(code) or next_code("ops_regions", "REG"), safe(name), safe(zone_level), float(q2(allowance_amount)), safe(notes), active_flag))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/regions?notice={notice}", status_code=303)


@router.get("/ui/operations/service-catalog")
def service_catalog_redirect():
    return RedirectResponse("/ui/operations/action-catalog", status_code=302)


@router.get("/ui/operations/action-catalog", response_class=HTMLResponse)
def service_catalog_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ops_service_catalog ORDER BY id DESC").fetchall()
    edit_row = conn.execute("SELECT * FROM ops_service_catalog WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_action_code(),
        "name": "",
        "service_category": "field_maintenance",
        "unit_price": "0.00",
        "technician_incentive": "0.00",
        "default_region_level": "zone_1",
        "default_duration_hours": "0.00",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        materials_count = conn.execute("SELECT COUNT(*) AS cnt FROM ops_action_materials WHERE action_id = ?", (row["id"],)).fetchone()["cnt"]
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['name'])}</td>
            <td>{safe(row['service_category']).replace('_', ' ').title()}</td>
            <td>{money(row['unit_price'])}</td>
            <td>{money(row['technician_incentive'])}</td>
            <td>{safe(row['default_region_level']).replace('_', ' ').title()}</td>
            <td>{safe(row['default_duration_hours'])}</td>
            <td>{materials_count}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td>
                <a class="btn blue" href="/ui/operations/action-catalog?edit_id={row['id']}">Edit</a>
                <a class="btn gray" href="/ui/operations/action-catalog/{row['id']}/materials">Materials</a>
            </td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='10' style='text-align:center;'>No actions added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Action Catalog</h2>
                <div class="section-note">Priced actions used in field maintenance, workshop repairs, and cabinet assembly. Each action can have a fixed list of raw materials.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn blue" href="/ui/operations/action-catalog/template.csv">Template CSV</a>
                <a class="btn blue" href="/ui/operations/action-catalog/template.xlsx">Template Excel</a>
                <a class="btn gray" href="/ui/operations">Back to Operations</a>
            </div>
        </div>
        <form method="post" action="/ui/operations/action-catalog/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Action Name</label><input name="name" value="{safe(form_values.get('name'))}" required></div>
                <div class="form-group"><label>Category</label><select name="service_category">{service_category_options(form_values.get('service_category'))}</select></div>
                <div class="form-group"><label>Unit Price</label><input name="unit_price" value="{safe(form_values.get('unit_price'))}"></div>
                <div class="form-group"><label>Technician Incentive</label><input name="technician_incentive" value="{safe(form_values.get('technician_incentive'))}"></div>
                <div class="form-group"><label>Default Region Level</label><select name="default_region_level">{zone_level_options(form_values.get('default_region_level'))}</select></div>
                <div class="form-group"><label>Default Duration (Hours)</label><input name="default_duration_hours" value="{safe(form_values.get('default_duration_hours'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Service</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Action" if safe_int(form_values.get('id')) > 0 else "Save Action"}</button>
                <a class="btn gray" href="/ui/operations/action-catalog">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <form method="post" action="/ui/operations/action-catalog/import" enctype="multipart/form-data">
            <div class="toolbar">
                <div>
                    <h3 class="sub-title" style="margin:0;">Import Actions</h3>
                    <div class="section-note">Use the template to import priced actions in bulk, then open each action and attach its fixed raw materials list.</div>
                </div>
            </div>
            <div class="form-grid" style="margin-top:14px;">
                <div class="form-group" style="grid-column: span 2;">
                    <label>Import File (CSV / XLSX)</label>
                    <input type="file" name="file" accept=".csv,.xlsx" required>
                </div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">Import Actions</button>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Action List</h3>
        <table>
            <tr><th>Code</th><th>Action</th><th>Category</th><th>Price</th><th>Incentive</th><th>Region Level</th><th>Hours</th><th>Materials</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    conn.close()
    return render_ops_page(request, "Action Catalog", html)


@router.post("/ui/operations/action-catalog/save")
def save_service(
    row_id: int = Form(0),
    code: str = Form(""),
    name: str = Form(""),
    service_category: str = Form("field_maintenance"),
    unit_price: str = Form("0"),
    technician_incentive: str = Form("0"),
    default_region_level: str = Form("zone_1"),
    default_duration_hours: str = Form("0"),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    if row_id > 0:
        conn.execute("""
            UPDATE ops_service_catalog
            SET code = ?, name = ?, service_category = ?, unit_price = ?, technician_incentive = ?, default_region_level = ?, default_duration_hours = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (
            safe(code),
            safe(name),
            safe(service_category),
            float(q2(unit_price)),
            float(q2(technician_incentive)),
            safe(default_region_level),
            float(q2(default_duration_hours)),
            safe(notes),
            active_flag,
            row_id,
        ))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_service_catalog (
                code, name, service_category, unit_price, technician_incentive,
                default_region_level, default_duration_hours, notes, is_active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            safe(code) or next_action_code(),
            safe(name),
            safe(service_category),
            float(q2(unit_price)),
            float(q2(technician_incentive)),
            safe(default_region_level),
            float(q2(default_duration_hours)),
            safe(notes),
            active_flag,
        ))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/action-catalog?notice={notice}", status_code=303)


@router.get("/ui/operations/action-catalog/template.csv")
def actions_template_csv():
    headers = [
        "code",
        "name",
        "service_category",
        "unit_price",
        "technician_incentive",
        "default_region_level",
        "default_duration_hours",
        "notes",
        "is_active",
    ]
    rows = [
        ["ACT-0001", "Replace Rectifier Fan", "workshop_repair", "1500", "250", "zone_1", "2", "Workshop repair action", "1"],
        ["ACT-0002", "Cabinet Assembly Visit", "cabinet_assembly", "3200", "450", "zone_2", "6", "Assembly and field finishing", "1"],
        ["ACT-0003", "BTS Power Fault Fix", "field_maintenance", "2800", "400", "zone_3", "4", "Standard field maintenance action", "1"],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    data = output.getvalue().encode("utf-8-sig")
    stream = io.BytesIO(data)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="action_catalog_template.csv"'},
    )


@router.get("/ui/operations/action-catalog/template.xlsx")
def actions_template_xlsx():
    if Workbook is None:
        return RedirectResponse("/ui/operations/action-catalog?notice=" + quote("Excel template is not available. Use CSV template."), status_code=302)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Action Catalog"
    sheet.append([
        "code",
        "name",
        "service_category",
        "unit_price",
        "technician_incentive",
        "default_region_level",
        "default_duration_hours",
        "notes",
        "is_active",
    ])
    sheet.append(["ACT-0001", "Replace Rectifier Fan", "workshop_repair", "1500", "250", "zone_1", "2", "Workshop repair action", "1"])
    sheet.append(["ACT-0002", "Cabinet Assembly Visit", "cabinet_assembly", "3200", "450", "zone_2", "6", "Assembly and field finishing", "1"])
    sheet.append(["ACT-0003", "BTS Power Fault Fix", "field_maintenance", "2800", "400", "zone_3", "4", "Standard field maintenance action", "1"])
    lookups = workbook.create_sheet("Lookups")
    lookups.append(["service_category", "region_level"])
    for idx in range(max(len(SERVICE_CATEGORIES), len(ZONE_LEVELS))):
        cat = SERVICE_CATEGORIES[idx][0] if idx < len(SERVICE_CATEGORIES) else ""
        zone = ZONE_LEVELS[idx][0] if idx < len(ZONE_LEVELS) else ""
        lookups.append([cat, zone])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="action_catalog_template.xlsx"'},
    )


@router.post("/ui/operations/action-catalog/import")
async def import_actions(file: UploadFile = File(...)):
    filename = safe(file.filename).lower()
    file_bytes = await file.read()
    if not filename:
        return RedirectResponse("/ui/operations/action-catalog?notice=" + quote("Please choose a file to import."), status_code=302)
    try:
        if filename.endswith(".csv"):
            raw_rows = parse_csv_rows(file_bytes)
        elif filename.endswith(".xlsx"):
            raw_rows = parse_xlsx_rows(file_bytes)
        else:
            return RedirectResponse("/ui/operations/action-catalog?notice=" + quote("Only CSV or XLSX files are supported."), status_code=302)
    except Exception as ex:
        return RedirectResponse("/ui/operations/action-catalog?notice=" + quote(f"Import failed: {safe(ex)}"), status_code=302)
    conn = get_conn()
    imported = 0
    updated = 0
    skipped = 0
    try:
        for raw_row in raw_rows:
            row = normalize_action_import_row(raw_row)
            if not safe(row["name"]):
                skipped += 1
                continue
            code = safe(row["code"]) or next_action_code()
            existing = conn.execute("SELECT id FROM ops_service_catalog WHERE UPPER(COALESCE(code,'')) = UPPER(?) LIMIT 1", (code,)).fetchone()
            payload = (
                code,
                safe(row["name"]),
                safe(row["service_category"]) or "field_maintenance",
                float(q2(row["unit_price"])),
                float(q2(row["technician_incentive"])),
                safe(row["default_region_level"]) or "zone_1",
                float(q2(row["default_duration_hours"])),
                safe(row["notes"]),
                1 if safe(row["is_active"]) not in ["0", "false", "False"] else 0,
            )
            if existing:
                conn.execute("""
                    UPDATE ops_service_catalog
                    SET code = ?, name = ?, service_category = ?, unit_price = ?, technician_incentive = ?,
                        default_region_level = ?, default_duration_hours = ?, notes = ?, is_active = ?
                    WHERE id = ?
                """, payload + (existing["id"],))
                updated += 1
            else:
                conn.execute("""
                    INSERT INTO ops_service_catalog (
                        code, name, service_category, unit_price, technician_incentive,
                        default_region_level, default_duration_hours, notes, is_active
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, payload)
                imported += 1
        conn.commit()
    except Exception as ex:
        conn.rollback()
        conn.close()
        return RedirectResponse("/ui/operations/action-catalog?notice=" + quote(f"Import failed: {safe(ex)}"), status_code=302)
    conn.close()
    msg = f"Import completed. Added: {imported}, Updated: {updated}, Skipped: {skipped}."
    return RedirectResponse("/ui/operations/action-catalog?notice=" + quote(msg), status_code=302)


@router.get("/ui/operations/action-catalog/{action_id}/materials", response_class=HTMLResponse)
def action_materials_page(request: Request, action_id: int):
    conn = get_conn()
    action = conn.execute("SELECT * FROM ops_service_catalog WHERE id = ? LIMIT 1", (action_id,)).fetchone()
    if not action:
        conn.close()
        return HTMLResponse("Action not found.", status_code=404)
    materials = get_action_materials(conn, action_id)
    conn.close()
    rows_html = ""
    total_cost = Decimal("0.00")
    for row in materials:
        line_total = q2(row["qty"]) * q2(row["unit_cost"])
        total_cost += line_total
        rows_html += f"""
        <tr>
            <td>{safe(row['line_no'])}</td>
            <td>{safe(row['item_code'])}</td>
            <td>{safe(row['item_name'])}</td>
            <td>{safe(row['uom'])}</td>
            <td>{money(row['qty'])}</td>
            <td>{money(row['unit_cost'])}</td>
            <td>{money(line_total)}</td>
            <td>{safe(row['notes'])}</td>
            <td>
                <form method="post" action="/ui/operations/action-catalog/{action_id}/materials/{row['id']}/delete" style="display:inline;">
                    <button class="btn red" type="submit">Delete</button>
                </form>
            </td>
        </tr>
        """
    if not rows_html:
        rows_html = "<tr><td colspan='9' style='text-align:center;'>No fixed raw materials added for this action.</td></tr>"
    html = f"""
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Action Materials</h2>
                <div class="section-note"><b>{safe(action['code'])}</b> - {safe(action['name'])}. These are the fixed raw materials used to cost the work order automatically.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn blue" href="/ui/operations/action-catalog?edit_id={action_id}">Back to Action</a>
                <a class="btn gray" href="/ui/operations/action-catalog">Back to Catalog</a>
            </div>
        </div>
        <div class="table-summary">
            <span class="summary-pill">Action Price: {money(action['unit_price'])}</span>
            <span class="summary-pill">Incentive: {money(action['technician_incentive'])}</span>
            <span class="summary-pill">Fixed Material Cost: {money(total_cost)}</span>
        </div>
        <form method="post" action="/ui/operations/action-catalog/{action_id}/materials/add" style="margin-top:16px;">
            <div class="form-grid">
                <div class="form-group" style="grid-column: span 2;"><label>Item</label><select name="item_id" required>{item_options()}</select></div>
                <div class="form-group"><label>Fixed Qty</label><input name="qty" value="1" required></div>
                <div class="form-group"><label>Standard Unit Cost</label><input name="unit_cost" value="0"></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">Add Fixed Material</button>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Fixed Raw Materials</h3>
        <table>
            <tr><th>#</th><th>Item Code</th><th>Item Name</th><th>UOM</th><th>Qty</th><th>Unit Cost</th><th>Line Cost</th><th>Notes</th><th>Action</th></tr>
            {rows_html}
        </table>
    </div>
    """
    return render_ops_page(request, "Action Materials", html)


@router.post("/ui/operations/action-catalog/{action_id}/materials/add")
def add_action_material(
    action_id: int,
    item_id: int = Form(...),
    qty: str = Form("1"),
    unit_cost: str = Form("0"),
    notes: str = Form(""),
):
    item = get_item_by_id(item_id)
    if not item:
        return HTMLResponse("Item not found.", status_code=404)
    resolved_cost = q2(unit_cost)
    if resolved_cost <= Decimal("0.00"):
        resolved_cost = get_item_last_cost(item_id)
    conn = get_conn()
    next_line = conn.execute("SELECT COALESCE(MAX(line_no), 0) + 1 AS next_line FROM ops_action_materials WHERE action_id = ?", (action_id,)).fetchone()["next_line"]
    conn.execute("""
        INSERT INTO ops_action_materials (
            action_id, line_no, item_id, item_code, item_name, uom, qty, unit_cost, notes
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        action_id,
        int(next_line or 1),
        item_id,
        safe(item["code"]),
        safe(item["name"]),
        safe(item["uom"]),
        float(q2(qty)),
        float(resolved_cost),
        safe(notes),
    ))
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/action-catalog/{action_id}/materials", status_code=303)


@router.post("/ui/operations/action-catalog/{action_id}/materials/{material_id}/delete")
def delete_action_material(action_id: int, material_id: int):
    conn = get_conn()
    conn.execute("DELETE FROM ops_action_materials WHERE id = ? AND action_id = ?", (material_id, action_id))
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/action-catalog/{action_id}/materials", status_code=303)


@router.get("/ui/operations/vehicle-rates", response_class=HTMLResponse)
def vehicle_rates_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("""
        SELECT vr.*, rs.name AS supplier_name
        FROM ops_vehicle_rates vr
        LEFT JOIN ops_rental_suppliers rs ON rs.id = vr.rental_supplier_id
        ORDER BY vr.id DESC
    """).fetchall()
    edit_row = conn.execute("SELECT * FROM ops_vehicle_rates WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_code("ops_vehicle_rates", "VEH"),
        "rental_supplier_id": "",
        "slab_name": "standard",
        "vehicle_type": "",
        "ticket_open_price": "0.00",
        "second_slab_to_km": "300.00",
        "km_rate_101_300": "0.00",
        "km_rate_over_300": "0.00",
        "waiting_hour_rate": "0.00",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['supplier_name']) or '-'}</td>
            <td>{safe(row['vehicle_type'])}</td>
            <td>{safe(row['slab_name']) or 'standard'}</td>
            <td>{money(row['ticket_open_price'])}</td>
            <td>{money(row['second_slab_to_km'])}</td>
            <td>{money(row['km_rate_101_300'])}</td>
            <td>{money(row['km_rate_over_300'])}</td>
            <td>{money(row['waiting_hour_rate'])}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td><a class="btn blue" href="/ui/operations/vehicle-rates?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='11' style='text-align:center;'>No vehicle rates added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Vehicle Rates</h2>
                <div class="section-note">Trip pricing is now flexible: fixed ticket up to 100 KM, then a second slab up to the limit you define here, then a final over-limit per-KM rate, plus waiting hours.</div>
            </div>
            <a class="btn gray" href="/ui/operations">Back to Operations</a>
        </div>
        <form method="post" action="/ui/operations/vehicle-rates/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Rental Office</label><select name="rental_supplier_id" required>{rental_supplier_options(form_values.get('rental_supplier_id'))}</select></div>
                <div class="form-group"><label>Vehicle Type</label><input name="vehicle_type" value="{safe(form_values.get('vehicle_type'))}" required></div>
                <div class="form-group"><label>Pricing Slab</label><input name="slab_name" value="{safe(form_values.get('slab_name') or 'standard')}" required></div>
                <div class="form-group"><label>0 - 100 KM Ticket Price</label><input name="ticket_open_price" value="{safe(form_values.get('ticket_open_price'))}"></div>
                <div class="form-group"><label>Second Slab Upper KM</label><input name="second_slab_to_km" value="{safe(form_values.get('second_slab_to_km', 300))}"></div>
                <div class="form-group"><label>101 - Upper Slab Rate</label><input name="km_rate_101_300" value="{safe(form_values.get('km_rate_101_300'))}"></div>
                <div class="form-group"><label>Over Upper Slab Rate</label><input name="km_rate_over_300" value="{safe(form_values.get('km_rate_over_300'))}"></div>
                <div class="form-group"><label>Waiting Hour Rate</label><input name="waiting_hour_rate" value="{safe(form_values.get('waiting_hour_rate'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Vehicle Rate</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="card" style="margin-top:14px;background:#f8fbff;">
                <div class="section-note">
                    Costing concept:
                    1. If trip KM <= 100: total = ticket price + waiting.
                    2. If trip KM is from 101 to the upper slab KM: total = ticket price for first 100 + ((KM - 100) x slab rate) + waiting.
                    3. If trip KM is above the upper slab KM: total = KM x over-slab rate + waiting.
                </div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Vehicle Rate" if safe_int(form_values.get('id')) > 0 else "Save Vehicle Rate"}</button>
                <a class="btn gray" href="/ui/operations/vehicle-rates">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Vehicle Rate List</h3>
        <table>
            <tr><th>Code</th><th>Rental Office</th><th>Vehicle Type</th><th>Pricing Slab</th><th>0-100 Ticket</th><th>Upper Slab KM</th><th>101-Upper Rate</th><th>Over Upper Rate</th><th>Waiting Hour</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Vehicle Rates", html)


@router.post("/ui/operations/vehicle-rates/save")
def save_vehicle_rate(
    row_id: int = Form(0),
    code: str = Form(""),
    rental_supplier_id: int = Form(0),
    slab_name: str = Form("standard"),
    vehicle_type: str = Form(""),
    ticket_open_price: str = Form("0"),
    second_slab_to_km: str = Form("300"),
    km_rate_101_300: str = Form("0"),
    km_rate_over_300: str = Form("0"),
    waiting_hour_rate: str = Form("0"),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    if safe_int(rental_supplier_id) <= 0:
        return RedirectResponse("/ui/operations/vehicle-rates?notice=" + quote("Please select the rental office for this pricing slab."), status_code=302)
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    if row_id > 0:
        conn.execute("""
            UPDATE ops_vehicle_rates
            SET code = ?, rental_supplier_id = ?, slab_name = ?, vehicle_type = ?, ticket_open_price = ?, second_slab_to_km = ?, km_rate_101_300 = ?, km_rate_over_300 = ?, km_rate = ?, waiting_hour_rate = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (
            safe(code),
            safe_int(rental_supplier_id),
            safe(slab_name) or "standard",
            safe(vehicle_type),
            float(q2(ticket_open_price)),
            float(q2(second_slab_to_km)),
            float(q2(km_rate_101_300)),
            float(q2(km_rate_over_300)),
            float(q2(km_rate_over_300)),
            float(q2(waiting_hour_rate)),
            safe(notes),
            active_flag,
            row_id,
        ))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_vehicle_rates (
                code, rental_supplier_id, slab_name, vehicle_type, ticket_open_price, second_slab_to_km, km_rate_101_300, km_rate_over_300, km_rate, waiting_hour_rate, notes, is_active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            safe(code) or next_code("ops_vehicle_rates", "VEH"),
            safe_int(rental_supplier_id),
            safe(slab_name) or "standard",
            safe(vehicle_type),
            float(q2(ticket_open_price)),
            float(q2(second_slab_to_km)),
            float(q2(km_rate_101_300)),
            float(q2(km_rate_over_300)),
            float(q2(km_rate_over_300)),
            float(q2(waiting_hour_rate)),
            safe(notes),
            active_flag,
        ))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/vehicle-rates?notice={notice}", status_code=303)


@router.get("/ui/operations/rental-suppliers", response_class=HTMLResponse)
def rental_suppliers_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ops_rental_suppliers ORDER BY id DESC").fetchall()
    edit_row = conn.execute("SELECT * FROM ops_rental_suppliers WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_rental_supplier_code(),
        "name": "",
        "contact_person": "",
        "phone": "",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['name'])}</td>
            <td>{safe(row['contact_person'])}</td>
            <td>{safe(row['phone'])}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td><a class="btn blue" href="/ui/operations/rental-suppliers?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='6' style='text-align:center;'>No rental offices added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Rental Offices</h2>
                <div class="section-note">Car rental suppliers or offices used for trip tickets. The same vehicle stays tied to one office.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn blue" href="/ui/operations/vehicles">Vehicles</a>
                <a class="btn gray" href="/ui/operations">Back to Operations</a>
            </div>
        </div>
        <form method="post" action="/ui/operations/rental-suppliers/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Office Name</label><input name="name" value="{safe(form_values.get('name'))}" required></div>
                <div class="form-group"><label>Contact Person</label><input name="contact_person" value="{safe(form_values.get('contact_person'))}"></div>
                <div class="form-group"><label>Phone</label><input name="phone" value="{safe(form_values.get('phone'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Rental Office</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Office" if safe_int(form_values.get('id')) > 0 else "Save Office"}</button>
                <a class="btn gray" href="/ui/operations/rental-suppliers">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Rental Office List</h3>
        <table>
            <tr><th>Code</th><th>Name</th><th>Contact</th><th>Phone</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Rental Offices", html)


@router.post("/ui/operations/rental-suppliers/save")
def save_rental_supplier(
    row_id: int = Form(0),
    code: str = Form(""),
    name: str = Form(""),
    contact_person: str = Form(""),
    phone: str = Form(""),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    if row_id > 0:
        conn.execute("""
            UPDATE ops_rental_suppliers
            SET code = ?, name = ?, contact_person = ?, phone = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, (safe(code), safe(name), safe(contact_person), safe(phone), safe(notes), active_flag, row_id))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_rental_suppliers (code, name, contact_person, phone, notes, is_active)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (safe(code) or next_rental_supplier_code(), safe(name), safe(contact_person), safe(phone), safe(notes), active_flag))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/rental-suppliers?notice={notice}", status_code=303)


@router.get("/ui/operations/vehicles", response_class=HTMLResponse)
def vehicles_page(request: Request, edit_id: int = 0, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            v.*,
            rs.name AS supplier_name,
            vr.code AS rate_code,
            vr.vehicle_type AS rate_vehicle_type,
            vr.slab_name AS rate_slab_name
        FROM ops_vehicles v
        LEFT JOIN ops_rental_suppliers rs ON rs.id = v.rental_supplier_id
        LEFT JOIN ops_vehicle_rates vr ON vr.id = v.vehicle_rate_id
        ORDER BY v.id DESC
    """).fetchall()
    edit_row = conn.execute("SELECT * FROM ops_vehicles WHERE id = ? LIMIT 1", (edit_id,)).fetchone() if edit_id else None
    conn.close()
    form_values = dict(edit_row) if edit_row else {
        "id": 0,
        "code": next_vehicle_code(),
        "vehicle_name": "",
        "vehicle_type": "",
        "rental_supplier_id": "",
        "vehicle_rate_id": "",
        "pricing_slab_name": "standard",
        "plate_no": "",
        "driver_source": "company_driver",
        "supplier_driver_name": "",
        "notes": "",
        "is_active": 1,
    }
    body = ""
    for row in rows:
        rate_label = safe(row["pricing_slab_name"]) or safe(row["rate_slab_name"]) or "standard"
        driver_label = "Company Driver" if safe(row["driver_source"]) == "company_driver" else "Supplier Driver"
        body += f"""
        <tr>
            <td>{safe(row['code'])}</td>
            <td>{safe(row['vehicle_name'])}</td>
            <td>{safe(row['vehicle_type'])}</td>
            <td>{safe(row['supplier_name'])}</td>
            <td>{safe(row['plate_no'])}</td>
            <td>{driver_label}</td>
            <td>{rate_label}</td>
            <td>{"Active" if int(row['is_active'] or 0) == 1 else "Inactive"}</td>
            <td><a class="btn blue" href="/ui/operations/vehicles?edit_id={row['id']}">Edit</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='9' style='text-align:center;'>No vehicles added yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Vehicles</h2>
                <div class="section-note">Each vehicle belongs to one rental office and one pricing slab. New trips automatically pick the latest active pricing slab for that office, type, and slab.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn blue" href="/ui/operations/rental-suppliers">Rental Offices</a>
                <a class="btn blue" href="/ui/operations/vehicle-rates">Vehicle Rates</a>
                <a class="btn gray" href="/ui/operations">Back to Operations</a>
            </div>
        </div>
        <form method="post" action="/ui/operations/vehicles/save">
            <input type="hidden" name="row_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Vehicle Code</label><input name="code" value="{safe(form_values.get('code'))}" required></div>
                <div class="form-group"><label>Vehicle Name</label><input name="vehicle_name" value="{safe(form_values.get('vehicle_name'))}" placeholder="Optional display name"></div>
                <div class="form-group"><label>Vehicle Type</label><input name="vehicle_type" value="{safe(form_values.get('vehicle_type'))}" required></div>
                <div class="form-group"><label>Rental Office</label><select name="rental_supplier_id" required>{rental_supplier_options(form_values.get('rental_supplier_id'))}</select></div>
                <div class="form-group"><label>Pricing Slab</label><input name="pricing_slab_name" value="{safe(form_values.get('pricing_slab_name') or 'standard')}" required></div>
                <div class="form-group"><label>Reference Rate Card (Optional)</label><select name="vehicle_rate_id">{vehicle_rate_options(form_values.get('vehicle_rate_id'))}</select></div>
                <div class="form-group"><label>Plate No</label><input name="plate_no" value="{safe(form_values.get('plate_no'))}"></div>
                <div class="form-group"><label>Driver Source</label><select name="driver_source">{driver_source_options(form_values.get('driver_source'))}</select></div>
                <div class="form-group"><label>Default Supplier Driver Name</label><input name="supplier_driver_name" value="{safe(form_values.get('supplier_driver_name'))}"></div>
                <div class="form-group"><label>Active</label><div style="padding-top:10px;"><label><input type="checkbox" name="is_active" value="1" {active_checkbox(int(form_values.get('is_active', 1) or 0) == 1)}> Active Vehicle</label></div></div>
                <div class="form-group" style="grid-column: span 2;"><label>Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Vehicle" if safe_int(form_values.get('id')) > 0 else "Save Vehicle"}</button>
                <a class="btn gray" href="/ui/operations/vehicles">Clear</a>
            </div>
        </form>
    </div>
    <div class="card">
        <h3 class="sub-title">Vehicle List</h3>
        <table>
            <tr><th>Code</th><th>Name</th><th>Type</th><th>Rental Office</th><th>Plate</th><th>Driver Source</th><th>Pricing Slab</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Vehicles", html)


@router.post("/ui/operations/vehicles/save")
def save_vehicle(
    row_id: int = Form(0),
    code: str = Form(""),
    vehicle_name: str = Form(""),
    vehicle_type: str = Form(""),
    rental_supplier_id: int = Form(0),
    vehicle_rate_id: int = Form(0),
    pricing_slab_name: str = Form("standard"),
    plate_no: str = Form(""),
    driver_source: str = Form("company_driver"),
    supplier_driver_name: str = Form(""),
    notes: str = Form(""),
    is_active: str = Form(""),
):
    if safe_int(rental_supplier_id) <= 0:
        return RedirectResponse("/ui/operations/vehicles?notice=" + quote("Please select a rental office."), status_code=302)
    conn = get_conn()
    active_flag = 1 if safe(is_active) == "1" else 0
    values = (
        safe(code),
        safe(vehicle_name),
        safe(vehicle_type),
        safe_int(rental_supplier_id),
        safe_int(vehicle_rate_id),
        safe(pricing_slab_name) or "standard",
        safe(plate_no),
        safe(driver_source) or "company_driver",
        safe(supplier_driver_name),
        safe(notes),
        active_flag,
    )
    if row_id > 0:
        conn.execute("""
            UPDATE ops_vehicles
            SET code = ?, vehicle_name = ?, vehicle_type = ?, rental_supplier_id = ?, vehicle_rate_id = ?,
                pricing_slab_name = ?, plate_no = ?, driver_source = ?, supplier_driver_name = ?, notes = ?, is_active = ?
            WHERE id = ?
        """, values + (row_id,))
        notice = "updated"
    else:
        conn.execute("""
            INSERT INTO ops_vehicles (
                code, vehicle_name, vehicle_type, rental_supplier_id, vehicle_rate_id,
                pricing_slab_name, plate_no, driver_source, supplier_driver_name, notes, is_active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            safe(code) or next_vehicle_code(),
            safe(vehicle_name),
            safe(vehicle_type),
            safe_int(rental_supplier_id),
            safe_int(vehicle_rate_id),
            safe(pricing_slab_name) or "standard",
            safe(plate_no),
            safe(driver_source) or "company_driver",
            safe(supplier_driver_name),
            safe(notes),
            active_flag,
        ))
        notice = "saved"
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/vehicles?notice={notice}", status_code=303)


@router.get("/ui/operations/trips", response_class=HTMLResponse)
def trips_page(request: Request, notice: str = ""):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            t.*,
            v.code AS vehicle_code,
            v.plate_no,
            rs.name AS supplier_name,
            COALESCE(e.employee_name, e.name, e.full_name) AS driver_name,
            (SELECT COUNT(*) FROM ops_trip_work_orders l WHERE l.trip_id = t.id) AS work_order_count
        FROM ops_trip_tickets t
        LEFT JOIN ops_vehicles v ON v.id = t.vehicle_id
        LEFT JOIN ops_rental_suppliers rs ON rs.id = t.rental_supplier_id
        LEFT JOIN employees e ON e.id = t.driver_employee_id
        ORDER BY t.id DESC
    """).fetchall()
    conn.close()
    body = ""
    for row in rows:
        driver_label = safe(row["driver_name"]) or safe(row["supplier_driver_name"]) or "-"
        vehicle_label = safe(row["vehicle_code"]) or safe(row["vehicle_type"])
        if safe(row["supplier_name"]):
            vehicle_label = f"{vehicle_label} / {safe(row['supplier_name'])}"
        cost_label = "Pending Accounting Approval" if safe(row["status"]).lower() != "approved" else money(row["total_cost"])
        alloc_label = "Pending" if safe(row["status"]).lower() != "approved" else money(row["allocated_cost_per_work_order"])
        body += f"""
        <tr>
            <td>{safe(row['trip_no'])}</td>
            <td>{safe(row['trip_date'])}</td>
            <td>{vehicle_label}</td>
            <td>{driver_label}</td>
            <td>{safe(row['work_order_count'])}</td>
            <td>{money(row['total_km']) if q2(row['total_km']) > 0 else '-'}</td>
            <td>{cost_label}</td>
            <td>{alloc_label}</td>
            <td>{trip_status_chip(row['status'])}</td>
            <td><a class="btn blue" href="/ui/operations/trips/{row['id']}">Open</a></td>
        </tr>
        """
    if not body:
        body = "<tr><td colspan='10' style='text-align:center;'>No trip tickets yet.</td></tr>"
    html = f"""
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">Trip Tickets</h2>
                <div class="section-note">Draft by technical manager, completion by movement manager, and accounting approval before transport cost hits work orders.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn green" href="/ui/operations/trips/new">+ New Draft Trip</a>
                <a class="btn blue" href="/ui/operations/vehicles">Vehicles</a>
                <a class="btn blue" href="/ui/operations/rental-suppliers">Rental Offices</a>
                <a class="btn gray" href="/ui/operations">Back to Operations</a>
            </div>
        </div>
        <div class="table-summary">
            <span class="summary-pill">Draft: {sum(1 for r in rows if safe(r['status']).lower() == 'draft')}</span>
            <span class="summary-pill">Dispatched: {sum(1 for r in rows if safe(r['status']).lower() == 'dispatched')}</span>
            <span class="summary-pill">Completed: {sum(1 for r in rows if safe(r['status']).lower() == 'completed')}</span>
            <span class="summary-pill">Approved: {sum(1 for r in rows if safe(r['status']).lower() == 'approved')}</span>
        </div>
    </div>
    <div class="card">
        <h3 class="sub-title">Trip Register</h3>
        <table>
            <tr><th>Trip No</th><th>Date</th><th>Vehicle</th><th>Driver</th><th>Work Orders</th><th>KM</th><th>Total Cost</th><th>Cost / WO</th><th>Status</th><th>Action</th></tr>
            {body}
        </table>
    </div>
    """
    return render_ops_page(request, "Trip Tickets", html)


@router.get("/ui/operations/trips/new", response_class=HTMLResponse)
def new_trip_page(request: Request, notice: str = ""):
    return trip_detail_page(request, 0, notice)


@router.get("/ui/operations/trips/{trip_id}", response_class=HTMLResponse)
def trip_detail_page(request: Request, trip_id: int, notice: str = ""):
    conn = get_conn()
    trip = None
    lines = []
    if trip_id > 0:
        trip = conn.execute("""
            SELECT
                t.*,
                v.code AS vehicle_code,
                v.plate_no,
                rs.name AS supplier_name,
                COALESCE(e.employee_name, e.name, e.full_name) AS driver_name
            FROM ops_trip_tickets t
            LEFT JOIN ops_vehicles v ON v.id = t.vehicle_id
            LEFT JOIN ops_rental_suppliers rs ON rs.id = t.rental_supplier_id
            LEFT JOIN employees e ON e.id = t.driver_employee_id
            WHERE t.id = ?
            LIMIT 1
        """, (trip_id,)).fetchone()
        if not trip:
            conn.close()
            return HTMLResponse("Trip not found.", status_code=404)
        lines = conn.execute("""
            SELECT
                l.*,
                wo.work_order_no,
                wo.site_name,
                wo.status,
                c.name AS company_name
            FROM ops_trip_work_orders l
            LEFT JOIN ops_work_orders wo ON wo.id = l.work_order_id
            LEFT JOIN ops_contract_companies c ON c.id = wo.company_id
            WHERE l.trip_id = ?
            ORDER BY l.line_no, l.id
        """, (trip_id,)).fetchall()
    conn.close()

    selected_work_orders = [row["work_order_id"] for row in lines]
    form_values = dict(trip) if trip else {
        "id": 0,
        "trip_no": next_trip_no(),
        "trip_date": "",
        "vehicle_id": "",
        "driver_employee_id": "",
        "supplier_driver_name": "",
        "notes": "",
        "status": "draft",
        "start_odometer": "",
        "end_odometer": "",
        "waiting_hours": "0",
        "movement_notes": "",
        "accounting_notes": "",
        "driver_source": "company_driver",
    }
    if trip and not form_values.get("supplier_driver_name") and safe(form_values.get("driver_source")) == "supplier_driver":
        form_values["supplier_driver_name"] = safe(form_values.get("driver_name"))

    line_rows = ""
    for idx, row in enumerate(lines, start=1):
        alloc = "Pending approval" if safe(form_values.get("status")).lower() != "approved" else money(row["allocated_cost"])
        line_rows += f"""
        <tr>
            <td>{idx}</td>
            <td>{safe(row['work_order_no'])}</td>
            <td>{safe(row['company_name'])}</td>
            <td>{safe(row['site_name'])}</td>
            <td>{safe(row['status']).title()}</td>
            <td>{alloc}</td>
        </tr>
        """
    if not line_rows:
        line_rows = "<tr><td colspan='6' style='text-align:center;'>No work orders linked yet.</td></tr>"

    status_value = safe(form_values.get("status") or "draft").lower()
    start_photo_html = f"<a class='btn gray' target='_blank' href='{safe(form_values.get('start_photo_path'))}'>View Start Meter</a>" if safe(form_values.get("start_photo_path")) else "<span class='section-note'>No start photo yet.</span>"
    end_photo_html = f"<a class='btn gray' target='_blank' href='{safe(form_values.get('end_photo_path'))}'>View End Meter</a>" if safe(form_values.get("end_photo_path")) else "<span class='section-note'>No end photo yet.</span>"
    driver_source_value = safe(form_values.get("driver_source") or "company_driver")
    trip_title = "New Draft Trip" if safe_int(form_values.get("id")) == 0 else f"Trip {safe(form_values.get('trip_no'))}"

    html = f"""
    <style>
        .checkbox-grid {{ display:grid; grid-template-columns: repeat(auto-fit, minmax(260px, 1fr)); gap:10px; margin-top:10px; }}
        .checkbox-card {{ display:flex; gap:10px; align-items:flex-start; padding:12px; border:1px solid #d6e1f4; border-radius:14px; background:#fff; }}
        .checkbox-card input {{ margin-top:4px; }}
        .empty-note {{ padding:14px; border:1px dashed #c9d8f0; border-radius:14px; color:#61779b; background:#f8fbff; }}
    </style>
    {current_form_notice(notice)}
    <div class="card">
        <div class="toolbar">
            <div>
                <h2 style="margin:0;">{trip_title}</h2>
                <div class="section-note">Draft trip is prepared by the technical manager. Vehicle can still be changed while the trip is in draft only.</div>
            </div>
            <div style="display:flex;gap:8px;flex-wrap:wrap;">
                <a class="btn blue" href="/ui/operations/trips">Back to Trips</a>
                <a class="btn gray" href="/ui/operations/vehicles">Vehicles</a>
            </div>
        </div>
        <div class="table-summary">
            <span class="summary-pill">Status: {safe(form_values.get('status')).replace('_', ' ').title()}</span>
            <span class="summary-pill">Trip No: {safe(form_values.get('trip_no'))}</span>
            <span class="summary-pill">Linked Work Orders: {len(lines) if lines else len(selected_work_orders)}</span>
            <span class="summary-pill">Total Cost: {"Pending Accounting Approval" if status_value != "approved" else money(form_values.get('total_cost'))}</span>
        </div>
    </div>
    <div class="card">
        <h3 class="sub-title">Draft Setup</h3>
        <form method="post" action="/ui/operations/trips/save-draft">
            <input type="hidden" name="trip_id" value="{safe(form_values.get('id', 0))}">
            <div class="form-grid">
                <div class="form-group"><label>Trip No</label><input name="trip_no" value="{safe(form_values.get('trip_no'))}" required></div>
                <div class="form-group"><label>Trip Date</label><input type="date" name="trip_date" value="{safe(form_values.get('trip_date'))}" required></div>
                <div class="form-group"><label>Vehicle</label><select name="vehicle_id" required {'' if status_value == 'draft' or safe_int(form_values.get('id')) == 0 else 'disabled'}>{vehicle_options(form_values.get('vehicle_id'))}</select></div>
                <div class="form-group"><label>Driver</label><select name="driver_employee_id">{employee_options(form_values.get('driver_employee_id'), driver_only=True)}</select></div>
                <div class="form-group"><label>Supplier Driver Name</label><input name="supplier_driver_name" value="{safe(form_values.get('supplier_driver_name'))}" placeholder="Only if the rental office provides the driver"></div>
                <div class="form-group" style="grid-column: span 2;"><label>Draft Notes</label><input name="notes" value="{safe(form_values.get('notes'))}"></div>
            </div>
            <div class="form-group">
                <label>Linked Work Orders</label>
                <div class="section-note">One trip can serve more than one work order, and later the trip cost will be divided equally between them after accounting approval.</div>
                <div class="checkbox-grid">
                    {work_order_choices(selected_work_orders)}
                </div>
            </div>
            <div class="form-actions">
                <button class="btn green" type="submit">{"Update Draft Trip" if safe_int(form_values.get('id')) > 0 else "Save Draft Trip"}</button>
                <a class="btn gray" href="/ui/operations/trips/new">Clear</a>
            </div>
        </form>
    </div>
    """

    if safe_int(form_values.get("id")) > 0:
        html += f"""
        <div class="card">
            <h3 class="sub-title">Linked Work Orders</h3>
            <table>
                <tr><th>#</th><th>Work Order</th><th>Company</th><th>Site</th><th>Status</th><th>Allocated Transport Cost</th></tr>
                {line_rows}
            </table>
        </div>
        <div class="card">
            <div class="toolbar">
                <div>
                    <h3 class="sub-title" style="margin:0;">Movement Manager Section</h3>
                    <div class="section-note">Start and end odometer cannot be saved without the meter photo. Waiting hours are entered manually.</div>
                </div>
                <div class="table-summary">
                    <span class="summary-pill">Start Photo: {"Saved" if safe(form_values.get('start_photo_path')) else "Missing"}</span>
                    <span class="summary-pill">End Photo: {"Saved" if safe(form_values.get('end_photo_path')) else "Missing"}</span>
                </div>
            </div>
            <form method="post" action="/ui/operations/trips/{trip_id}/movement" enctype="multipart/form-data">
                <div class="form-grid">
                    <div class="form-group"><label>Start Odometer</label><input name="start_odometer" value="{safe(form_values.get('start_odometer'))}"></div>
                    <div class="form-group"><label>Start Meter Photo</label><input type="file" name="start_photo" accept="image/*"></div>
                    <div class="form-group">{start_photo_html}</div>
                    <div class="form-group"><label>End Odometer</label><input name="end_odometer" value="{safe(form_values.get('end_odometer'))}"></div>
                    <div class="form-group"><label>End Meter Photo</label><input type="file" name="end_photo" accept="image/*"></div>
                    <div class="form-group">{end_photo_html}</div>
                    <div class="form-group"><label>Waiting Hours</label><input name="waiting_hours" value="{safe(form_values.get('waiting_hours'))}"></div>
                    <div class="form-group" style="grid-column: span 2;"><label>Movement Notes</label><input name="movement_notes" value="{safe(form_values.get('movement_notes'))}"></div>
                </div>
                <div class="form-actions">
                    <button class="btn blue" type="submit" name="action" value="dispatch" {"disabled" if status_value in ['completed', 'approved'] else ''}>Save Start / Dispatch</button>
                    <button class="btn green" type="submit" name="action" value="complete" {"disabled" if status_value == 'approved' else ''}>Complete Trip</button>
                </div>
            </form>
            <div class="table-summary">
                <span class="summary-pill">KM: {money(form_values.get('total_km')) if q2(form_values.get('total_km')) > 0 else '-'}</span>
                <span class="summary-pill">Trip Cost: {money(form_values.get('total_cost')) if q2(form_values.get('total_cost')) > 0 else '-'}</span>
                <span class="summary-pill">Driver Commission: {money(form_values.get('driver_commission_amount')) if q2(form_values.get('driver_commission_amount')) > 0 else '-'}</span>
            </div>
        </div>
        <div class="card">
            <div class="toolbar">
                <div>
                    <h3 class="sub-title" style="margin:0;">Accounting Approval</h3>
                    <div class="section-note">Trip cost should only become visible on work orders after accounting approval.</div>
                </div>
                <div class="table-summary">
                    <span class="summary-pill">Per Work Order: {money(form_values.get('allocated_cost_per_work_order')) if status_value == 'approved' else 'Pending'}</span>
                </div>
            </div>
            <form method="post" action="/ui/operations/trips/{trip_id}/approve">
                <div class="form-grid">
                    <div class="form-group" style="grid-column: span 2;"><label>Accounting Notes</label><input name="accounting_notes" value="{safe(form_values.get('accounting_notes'))}"></div>
                </div>
                <div class="form-actions">
                    <button class="btn green" type="submit" {"disabled" if status_value != 'completed' else ''}>Approve Trip</button>
                </div>
            </form>
        </div>
        """
    return render_ops_page(request, "Trip Ticket", html)


@router.post("/ui/operations/trips/save-draft")
async def save_trip_draft(
    request: Request,
    trip_id: int = Form(0),
    trip_no: str = Form(""),
    trip_date: str = Form(""),
    vehicle_id: int = Form(0),
    driver_employee_id: int = Form(0),
    supplier_driver_name: str = Form(""),
    notes: str = Form(""),
    work_order_ids: list[int] = Form([]),
):
    if not safe(trip_date):
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Trip date is required."), status_code=302)
    if safe_int(vehicle_id) <= 0:
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Please select the actual vehicle."), status_code=302)
    selected_work_orders = [safe_int(x) for x in work_order_ids if safe_int(x) > 0]
    if not selected_work_orders:
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Please link at least one work order to the trip."), status_code=302)

    vehicle = get_vehicle_row(vehicle_id)
    if not vehicle:
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Selected vehicle was not found."), status_code=302)
    if safe_int(vehicle["vehicle_rate_id_resolved"]) <= 0:
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("No active pricing slab was found for this vehicle. Please activate a matching vehicle pricing slab first."), status_code=302)
    driver_source = safe(vehicle["driver_source"]) or "company_driver"
    if driver_source == "company_driver" and safe_int(driver_employee_id) <= 0:
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Please select the company driver from HR employees."), status_code=302)
    if driver_source == "supplier_driver" and not safe(supplier_driver_name) and not safe(vehicle["supplier_driver_name"]):
        target = "/ui/operations/trips/new" if trip_id <= 0 else f"/ui/operations/trips/{trip_id}"
        return RedirectResponse(target + "?notice=" + quote("Please enter the supplier driver name for this trip."), status_code=302)

    conn = get_conn()
    current_trip = conn.execute("SELECT id, status FROM ops_trip_tickets WHERE id = ? LIMIT 1", (trip_id,)).fetchone() if trip_id > 0 else None
    if current_trip and safe(current_trip["status"]).lower() != "draft":
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Only draft trips can be edited by the technical manager."), status_code=302)

    base_values = (
        safe(trip_no) or next_trip_no(),
        safe(trip_date),
        safe_int(vehicle_id),
        safe_int(vehicle["rental_supplier_id"]),
        safe(vehicle["vehicle_type"]),
        safe_int(vehicle["vehicle_rate_id_resolved"]),
        driver_source,
        safe_int(driver_employee_id) if driver_source == "company_driver" else 0,
        safe(supplier_driver_name) or safe(vehicle["supplier_driver_name"]),
        float(q2(vehicle["ticket_open_price"] if vehicle else 0)),
        float(q2(vehicle["second_slab_to_km"] if vehicle else 300)),
        float(q2(vehicle["km_rate_101_300"] if vehicle else 0)),
        float(q2(vehicle["km_rate_over_300"] if vehicle else 0)),
        float(q2(vehicle["waiting_hour_rate"] if vehicle else 0)),
        safe(notes),
    )
    if trip_id > 0:
        conn.execute("""
            UPDATE ops_trip_tickets
            SET trip_no = ?, trip_date = ?, vehicle_id = ?, rental_supplier_id = ?, vehicle_type = ?,
                vehicle_rate_id = ?, driver_source = ?, driver_employee_id = ?, supplier_driver_name = ?,
                ticket_open_price = ?, second_slab_to_km = ?, km_rate_101_300 = ?, km_rate_over_300 = ?, waiting_hour_rate = ?, notes = ?
            WHERE id = ?
        """, base_values + (trip_id,))
        target_trip_id = trip_id
    else:
        conn.execute("""
            INSERT INTO ops_trip_tickets (
                trip_no, trip_date, vehicle_id, rental_supplier_id, vehicle_type, vehicle_rate_id,
                driver_source, driver_employee_id, supplier_driver_name,
                ticket_open_price, second_slab_to_km, km_rate_101_300, km_rate_over_300, waiting_hour_rate,
                notes, created_by
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, base_values + (actor_name_from_request(request),))
        target_trip_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

    conn.execute("DELETE FROM ops_trip_work_orders WHERE trip_id = ?", (target_trip_id,))
    for idx, work_order_id in enumerate(selected_work_orders, start=1):
        conn.execute("""
            INSERT OR IGNORE INTO ops_trip_work_orders (trip_id, work_order_id, line_no, allocated_cost)
            VALUES (?, ?, ?, 0)
        """, (target_trip_id, work_order_id, idx))
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/trips/{target_trip_id}?notice=" + quote("Draft trip saved successfully."), status_code=303)


@router.post("/ui/operations/trips/{trip_id}/movement")
async def save_trip_movement(
    request: Request,
    trip_id: int,
    action: str = Form("dispatch"),
    start_odometer: str = Form(""),
    end_odometer: str = Form(""),
    waiting_hours: str = Form("0"),
    movement_notes: str = Form(""),
    start_photo: UploadFile = File(None),
    end_photo: UploadFile = File(None),
):
    conn = get_conn()
    trip = conn.execute("SELECT * FROM ops_trip_tickets WHERE id = ? LIMIT 1", (trip_id,)).fetchone()
    if not trip:
        conn.close()
        return HTMLResponse("Trip not found.", status_code=404)
    if safe(trip["status"]).lower() == "approved":
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Approved trips cannot be changed."), status_code=302)

    start_value = q2(start_odometer) if safe(start_odometer) else q2(trip["start_odometer"])
    end_value = q2(end_odometer) if safe(end_odometer) else q2(trip["end_odometer"])
    waiting_value = q2(waiting_hours) if safe(waiting_hours) else q2(trip["waiting_hours"])
    start_photo_path = safe(trip["start_photo_path"])
    end_photo_path = safe(trip["end_photo_path"])
    if start_photo and safe(start_photo.filename):
        start_photo_path = await save_trip_photo(start_photo, f"trip_{trip_id}_start")
    if end_photo and safe(end_photo.filename):
        end_photo_path = await save_trip_photo(end_photo, f"trip_{trip_id}_end")

    if safe(action).lower() == "dispatch":
        if start_value <= Decimal("0.00"):
            conn.close()
            return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Start odometer is required before dispatch."), status_code=302)
        if not safe(start_photo_path):
            conn.close()
            return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Please upload the start meter photo before dispatch."), status_code=302)
        conn.execute("""
            UPDATE ops_trip_tickets
            SET start_odometer = ?, start_photo_path = ?, movement_notes = ?, status = 'dispatched'
            WHERE id = ?
        """, (float(start_value), safe(start_photo_path), safe(movement_notes), trip_id))
        conn.commit()
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Trip dispatched successfully."), status_code=303)

    if start_value <= Decimal("0.00"):
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Start odometer is required before completion."), status_code=302)
    if not safe(start_photo_path):
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Start meter photo is required before completion."), status_code=302)
    if end_value <= Decimal("0.00"):
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("End odometer is required to complete the trip."), status_code=302)
    if not safe(end_photo_path):
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Please upload the end meter photo before completing the trip."), status_code=302)
    if end_value < start_value:
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("End odometer cannot be less than start odometer."), status_code=302)

    total_km = q2(end_value - start_value)
    total_cost = calculate_trip_cost(
        total_km,
        waiting_value,
        trip["ticket_open_price"],
        trip["second_slab_to_km"],
        trip["km_rate_101_300"],
        trip["km_rate_over_300"],
        trip["waiting_hour_rate"],
    )
    work_order_count = conn.execute("SELECT COUNT(*) AS cnt FROM ops_trip_work_orders WHERE trip_id = ?", (trip_id,)).fetchone()["cnt"] or 0
    allocated_cost = q2(total_cost / Decimal(str(work_order_count))) if int(work_order_count) > 0 else Decimal("0.00")
    driver_commission_pct = Decimal("0.00")
    driver_commission_amount = Decimal("0.00")
    if safe(trip["driver_source"]) == "company_driver" and safe_int(trip["driver_employee_id"]) > 0:
        driver = conn.execute("SELECT COALESCE(trip_commission_pct, 0) AS trip_commission_pct FROM employees WHERE id = ? LIMIT 1", (trip["driver_employee_id"],)).fetchone()
        driver_commission_pct = q2(driver["trip_commission_pct"] if driver else 0)
        driver_commission_amount = q2(total_cost * driver_commission_pct / Decimal("100.00"))

    conn.execute("""
        UPDATE ops_trip_tickets
        SET start_odometer = ?, start_photo_path = ?, end_odometer = ?, end_photo_path = ?,
            waiting_hours = ?, total_km = ?, total_cost = ?, allocated_cost_per_work_order = ?,
            driver_commission_pct = ?, driver_commission_amount = ?, movement_notes = ?,
            movement_closed_by = ?, completed_at = CURRENT_TIMESTAMP, status = 'completed'
        WHERE id = ?
    """, (
        float(start_value),
        safe(start_photo_path),
        float(end_value),
        safe(end_photo_path),
        float(waiting_value),
        float(total_km),
        float(total_cost),
        float(allocated_cost),
        float(driver_commission_pct),
        float(driver_commission_amount),
        safe(movement_notes),
        actor_name_from_request(request),
        trip_id,
    ))
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Trip completed and ready for accounting approval."), status_code=303)


@router.post("/ui/operations/trips/{trip_id}/approve")
def approve_trip(
    request: Request,
    trip_id: int,
    accounting_notes: str = Form(""),
):
    conn = get_conn()
    trip = conn.execute("SELECT * FROM ops_trip_tickets WHERE id = ? LIMIT 1", (trip_id,)).fetchone()
    if not trip:
        conn.close()
        return HTMLResponse("Trip not found.", status_code=404)
    if safe(trip["status"]).lower() != "completed":
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Only completed trips can be approved by accounting."), status_code=302)
    lines = conn.execute("SELECT id FROM ops_trip_work_orders WHERE trip_id = ? ORDER BY id", (trip_id,)).fetchall()
    if not lines:
        conn.close()
        return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("No work orders linked to this trip."), status_code=302)
    allocated_cost = q2(trip["allocated_cost_per_work_order"])
    for row in lines:
        conn.execute("UPDATE ops_trip_work_orders SET allocated_cost = ? WHERE id = ?", (float(allocated_cost), row["id"]))
    conn.execute("""
        UPDATE ops_trip_tickets
        SET status = 'approved', accounting_notes = ?, approved_by = ?, approved_at = CURRENT_TIMESTAMP
        WHERE id = ?
    """, (safe(accounting_notes), actor_name_from_request(request), trip_id))
    conn.commit()
    conn.close()
    return RedirectResponse(f"/ui/operations/trips/{trip_id}?notice=" + quote("Trip approved successfully. Transport cost is now ready for work orders."), status_code=303)
